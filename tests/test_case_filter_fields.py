"""人员、关联字段及自定义日期的真实列表/导出回归，不执行台架。"""

import json
from io import BytesIO
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from test_case_governance import governance  # noqa: F401
from test_case_features import features  # noqa: F401
from models.test_case import CaseAttachment
from models.case_features import CaseTemplate
from models.filter_field import FilterField
from models.case_governance import CaseVersion
from services.case_query import query_cases, matches


def condition(field, operator, value=None):
    return {"field": field, "operator": operator, "value": value}


def list_ids(g, conditions, logic="and", **extra):
    response = g["client"].get(
        "/api/v1/test-cases",
        params={
            "project_id": g["project"].id,
            "filters": json.dumps({"conditions": conditions, "logic": logic}),
            **extra,
        },
    )
    assert response.status_code == 200, response.text
    return {v["id"] for v in response.json()["data"]["items"]}


def test_members_current_user_and_saved_token_are_dynamic(governance):
    g = governance
    first, second = g["cases"]
    second.created_by = g["users"][1].id
    second.updated_by = g["users"][1].id
    g["db"].commit()
    rows = [condition("createdBy", "belongs_to", ["CURRENT_USER"])]
    saved = g["client"].post(
        g["base"] + "/views",
        json={"name": "动态本人视图", "filters": {"filterConditions": rows}},
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["filters"]["filterConditions"] == rows
    assert list_ids(g, rows) == {first.id}
    assert list_ids(g, [condition("updatedBy", "belongs_to", [g["users"][1].id])]) == {
        second.id
    }
    assert list_ids(
        g, [condition("createdBy", "not_belongs_to", ["CURRENT_USER"])]
    ) == {second.id}
    g["state"]["user"] = g["users"][1]
    assert list_ids(g, rows) == {second.id}
    exported = g["client"].get(
        f"/api/v1/projects/{g['project'].id}/cases/export",
        params={"filters": json.dumps({"conditions": rows})},
    )
    assert exported.status_code == 403, exported.text
    g["state"]["user"] = g["users"][0]
    exported = g["client"].get(
        f"/api/v1/projects/{g['project'].id}/cases/export",
        params={"filters": json.dumps({"conditions": rows})},
    )
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(BytesIO(exported.content), read_only=True).active
    assert [r[0] for r in list(sheet.values)[1:]] == [first.case_code]
    assert rows[0]["value"] == ["CURRENT_USER"]
    g["state"]["user"] = g["users"][3]
    assert (
        g["client"]
        .get("/api/v1/test-cases", params={"project_id": g["project"].id})
        .status_code
        == 403
    )
    with pytest.raises(HTTPException) as exc:
        query_cases(g["db"], g["project"].id, filters={"conditions": rows})
    assert exc.value.status_code == 422


def test_attachment_any_match_negation_empty_and_export_match(features):
    g = features
    first, second = g["cases"]
    for name in ["报告验收.pdf", "说明.txt"]:
        response = g["client"].post(
            g["features"] + f"/cases/{first.id}/attachments",
            files={"file": (name, b"%PDF-sample", "application/pdf")},
        )
        assert response.status_code == 200, response.text
    # 同名的其他项目附件不能扩大当前项目结果。
    g["db"].add(
        CaseAttachment(
            case_id=g["foreign"].id,
            file_name="报告验收.pdf",
            file_path="未读取的隔离路径",
            file_size=1,
        )
    )
    g["db"].commit()
    versions = g["db"].query(CaseVersion).count()
    rows = [condition("attachment", "contains", "报告")]
    assert list_ids(g, rows) == {first.id}
    assert list_ids(g, [condition("attachment", "equals", "报告验收.pdf")]) == {
        first.id
    }
    assert list_ids(g, [condition("attachment", "not_contains", "报告")]) == {second.id}
    assert list_ids(g, [condition("attachment", "not_equals", "说明.txt")]) == {
        second.id
    }
    assert list_ids(g, [condition("attachment", "is_empty")]) == {second.id}
    assert list_ids(g, [condition("attachment", "is_not_empty")]) == {first.id}
    response = g["client"].get(
        f"/api/v1/projects/{g['project'].id}/cases/export",
        params={"filters": json.dumps({"conditions": rows})},
    )
    assert response.status_code == 200, response.text
    exported = list(
        load_workbook(BytesIO(response.content), read_only=True).active.values
    )
    assert [r[0] for r in exported[1:]] == [first.case_code]
    assert g["db"].query(CaseVersion).count() == versions


def test_linked_requirement_names_and_legacy_reference_are_searchable(features):
    g = features
    first, second = g["cases"]
    second.requirement_ref = "历史需求编号REQ-123"
    g["db"].commit()
    for kind in ["requirement", "defect"]:
        response = g["client"].post(
            g["features"] + "/issues",
            json={
                "kind": kind,
                "title": "登录需求" if kind == "requirement" else "排除的同名缺陷",
            },
        )
        assert response.status_code == 200, response.text
        assert (
            g["client"]
            .post(
                g["features"] + f"/cases/{first.id}/issues",
                json={"issueId": response.json()["data"]["id"]},
            )
            .status_code
            == 200
        )
    assert list_ids(g, [condition("requirementRef", "contains", "登录")]) == {first.id}
    assert list_ids(g, [condition("requirementRef", "contains", "REQ-123")]) == {
        second.id
    }
    assert list_ids(g, [condition("requirementRef", "contains", "排除")]) == set()
    assert list_ids(g, [condition("requirementRef", "is_not_empty")]) == {
        first.id,
        second.id,
    }
    assert list_ids(g, [condition("requirementRef", "not_contains", "登录")]) == {
        second.id
    }


def test_custom_date_all_comparisons_use_type_without_changing_text(governance):
    g = governance
    template = CaseTemplate(
        project_id=g["project"].id,
        created_by=g["users"][0].id,
        name="日期筛选模板",
        fields=[
            {"key": "deadline", "type": "date"},
            {"key": "literal", "type": "text"},
        ],
        defaults={},
    )
    g["db"].add(template)
    g["db"].flush()
    first, second = g["cases"]
    for case, day in [(first, "2026-10-05"), (second, "2026-10-06")]:
        case.template_id = template.id
        case.custom_fields = {"deadline": day, "literal": day}
    g["db"].commit()
    assert list_ids(
        g, [condition("customFields.deadline", "gt", "2026-10-05T00:00:00Z")]
    ) == {second.id}
    assert list_ids(
        g, [condition("customFields.deadline", "lt", "2026-10-06T00:00:00+08:00")]
    ) == {first.id}
    assert list_ids(
        g, [condition("customFields.deadline", "equals", "2026-10-05T08:00:00+08:00")]
    ) == {first.id}
    assert list_ids(
        g, [condition("customFields.deadline", "between", ["2026-10-05", "2026-10-05"])]
    ) == {first.id}
    assert list_ids(
        g, [condition("customFields.deadline", "gte", "2026-10-06")]
    ) == {second.id}
    assert list_ids(
        g, [condition("customFields.deadline", "lte", "2026-10-05")]
    ) == {first.id}
    assert list_ids(
        g, [condition("customFields.deadline", "not_equals", "2026-10-05")]
    ) == {second.id}
    assert (
        list_ids(
            g, [condition("customFields.literal", "equals", "2026-10-05T00:00:00Z")]
        )
        == set()
    )
    for field in ["createdAt", "updatedAt", "customFields.deadline"]:
        response = g["client"].get(
            "/api/v1/test-cases",
            params={
                "project_id": g["project"].id,
                "search": "不匹配任何用例",
                "filters": json.dumps(
                    {"conditions": [condition(field, "gt", "无效日期")]}
                ),
            },
        )
        assert response.status_code == 422, response.text
    assert first.custom_fields["deadline"] == "2026-10-05"


def test_project_field_upgrade_preserves_disabled_config_without_writes(governance):
    g = governance
    g["db"].add_all(
        [
            FilterField(
                project_id=g["project"].id,
                field_key="name",
                field_label="自定义名称",
                field_type="text",
                is_enabled=True,
            ),
            FilterField(
                project_id=g["project"].id,
                field_key="tags",
                field_label="标签",
                field_type="tags",
                is_enabled=False,
            ),
        ]
    )
    g["db"].commit()
    response = g["client"].get(
        "/api/v1/test-cases/filter-fields", params={"project_id": g["project"].id}
    )
    assert response.status_code == 200, response.text
    fields = {v["fieldKey"]: v for v in response.json()["data"]}
    assert fields["name"]["fieldLabel"] == "自定义名称" and "tags" not in fields
    assert (
        fields["createdBy"]["fieldType"] == "member"
        and fields["updatedBy"]["fieldType"] == "member"
    )
    assert fields["attachment"]["fieldType"] == "text"
    assert g["db"].query(FilterField).count() == 2


def test_multiple_tags_follow_official_any_positive_and_none_negative():
    assert matches(["甲"], "contains", ["甲", "乙"])
    assert not matches(["甲"], "not_contains", ["甲", "乙"])
    assert matches([], "not_contains", ["甲", "乙"])
