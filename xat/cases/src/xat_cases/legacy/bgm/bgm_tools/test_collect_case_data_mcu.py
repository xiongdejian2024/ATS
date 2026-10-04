import os
import sys
import re
from time import sleep
import pytest
import allure
import openpyxl

project_root = os.path.join(os.getcwd().split("sat")[0], 'sat')
sys.path.append(project_root)

work_path_2 = os.path.join(os.getcwd().split("sat")[0], 'sat', 'ecu_simulator', 'interface')
sys.path.append(work_path_2)

sys.path.append(os.getcwd())
sys.path.append(os.path.join(os.getcwd(), ".."))
sys.path.append(os.path.join(os.getcwd(), "../.."))
sys.path.append(os.path.join(os.getcwd(), "../../.."))

from xat_ecu.legacy.interface.ms.ms_lib import *
from xat_ecu.legacy.interface.ms.ms_to_jama import *

script_and_function_module_mapping = {
    "/BGM_MCU/MCU性能/业务性能/车辆公告发出时长": ["test_vehicleAnnouncement.py"],
    "/BGM_MCU/MCU基础通讯": ["test_switch.py"],
    "/BGM_MCU/MCU基础功能/外灯": [],
    "/BGM_MCU/MCU基础功能/VMM": [],
    "/BGM_MCU/MCU 基础平台/NVM存储": ["test_nvm_storage.py"],
    "/BGM_MCU/MCU 基础平台/信号路由": ["test_can_lin_fr_signal_route.py"],
    "/BGM_MCU/MCU 基础平台/MCU诊断路由": ["test_diag_route_app.py","test_diag_route_boot.py"],
    "/BGM_MCU/MCU 基础平台/CDD打包功能": ["test_cdd_function.py","test_cdd_performance.py"],
}



def dealwith_collect_orignal_data(file_path):
    pattern_0 = r"<Module\s(.*?)>"
    pattern_1 = r"_\d+"
    pattern_2 = r"\[(\d{5,})\]"
    poslist = []
    infodic = {}
    
    # r"C:\Users\hui.zhao\Desktop\Learning\Python\collect_case_info.log"
    with open(file_path,'r', encoding='utf-8') as f:  
        file_cont = str(f.readlines())   

    matches = re.findall(pattern_0, file_cont)
    for i in range(len(matches)):
        poslist.append(file_cont.find(matches[i]))
        if i >0 :
            if i < len(matches):
                str_1 = file_cont[poslist[i-1]:poslist[i]]
            else:
                str_1= file_cont[poslist[i-1]:len(file_cont)]
            matches_1 = re.findall(pattern_1, str_1)
            matches_2 = re.findall(pattern_2, str_1)
            for n in range(len(matches_1)):
                matches_1[n] = int(matches_1[n].replace("_",""))

            for m in range(len(matches_2)):
                matches_2[m] = int(matches_2[m].replace("case[",""))

            pos = matches[i-1].rfind("test_")
            infodic[matches[i-1][pos:]] = matches_2 + matches_1
    return infodic


def clear_excel_content(sheet_obj):
    for l in range(2,100):
        for r in range(1,20):
            sheet_obj.cell(row=l,column=r).value = ""

    
def dealwith_ms_original_case(original_data_list,tag):
    global script_and_function_module_mapping
    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    ms_data_dic_1 = {}
    for i in range(len(original_data_list)):
        # if "1.4" in original_data_list[i].get("tags"):
        #     print("V1.4")
        #     continue
        for n in range(len(all_function_module_name_list)):
            if tag == "smoke":
                tag_list = json.loads(original_data_list[i].get("tags"))
                for num in tag_list:
                    for elem in ["smoke","Smoke","smoking","Smoking"]:
                        if elem == num:
                            if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                                if all_function_module_name_list[n] in ms_data_dic_1.keys():
                                    if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                                        ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]  
            if tag == "sanity":
                tag_list = json.loads(original_data_list[i].get("tags"))
                for num in tag_list:
                    for elem in ["smoke","Smoke","smoking","Smoking","Checklist","checklist","sanity","Sanity"]:
                        if elem == num:
                            if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                                if all_function_module_name_list[n] in ms_data_dic_1.keys():
                                    if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                                        ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]
                                break
        
            if tag == "full":
                if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                    if all_function_module_name_list[n] in ms_data_dic_1.keys():
                        if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                            ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                    else:
                        ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]
    return ms_data_dic_1
    
    
def compare_ms_and_auto_case_info(ms_info,auto_script_info,excel_obj,tag):
    global cript_and_function_module_mapping
    
    if tag == "smoke":
        sht=excel_obj.get_sheet_by_name("smoke")
    elif tag == "sanity":
        sht=excel_obj.get_sheet_by_name("sanity")
    elif tag == "full":
        sht=excel_obj.get_sheet_by_name("full")
    
    clear_excel_content(sht)
        
    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    same_case_dic = {}
    coverage_rate_total = 0
    ms_case_num_total = 0
    auto_case_num_total = 0
    same_case_num_total = 0
    uncovernum_total = 0
 
    for i in range(len(all_function_module_name_list)):
        reason = []
        print(all_function_module_name_list[i])
        sht.cell(row=i+2,column=1).value = all_function_module_name_list[i]
        if all_function_module_name_list[i] not in list(ms_info.keys()):
            ms_case = []
        else:
            ms_case = ms_info[all_function_module_name_list[i]]

        ms_case_num = len(ms_case)
            
        related_script_file_name = script_and_function_module_mapping[all_function_module_name_list[i]]
        print(related_script_file_name)
        if type(related_script_file_name) == list:
            auto_case = []
            for elem in  related_script_file_name:
                if elem in list(auto_script_info.keys()):
                    auto_case = auto_case + auto_script_info[elem]
        else:
            if related_script_file_name not in list(auto_script_info.keys()):
                auto_case = []
            else:
                auto_case = auto_script_info[related_script_file_name]
        auto_case_num = len(auto_case)
            
        same_case_list = []
        for ms_case_id in ms_case:
            for auto_case_id in auto_case:
                if ms_case_id == auto_case_id and ms_case_id not in same_case_list:
                    same_case_list.append(ms_case_id)
        
        same_case_num = len(same_case_list)
        
        if ms_case_num != 0:
            coverage_rate = len(same_case_list)/ms_case_num    
            percentage = f"{coverage_rate*100:.1f}%"
        else:
            percentage = f"0%,可能此模块没有{tag}用例,也可能标签不正确，请进一步确认"
        
        
        
        sht.cell(row=i+2,column=2).value = str(percentage)
        auto_not_include = [x for x in ms_case if (x not in same_case_list)]
        uncovernum = len(auto_not_include)

        sht.cell(row=i+2,column=3).value = str(uncovernum)
        sht.cell(row=i+2,column=4).value = str(auto_not_include)
        sht.cell(row=i+2,column=5).value = str(ms_case_num)
        sht.cell(row=i+2,column=6).value = str(ms_case)
        sht.cell(row=i+2,column=7).value = str(auto_case_num)
        sht.cell(row=i+2,column=8).value = str(auto_case)
        sht.cell(row=i+2,column=9).value = str(len(same_case_list))
        sht.cell(row=i+2,column=10).value = str(same_case_list)
        
        ms_not_include = [x for x in ms_case if x not in same_case_list]
        sht.cell(row=i+2,column=11).value = str(ms_not_include)
        
        same_case_dic[all_function_module_name_list[i]] = same_case_list   
        sht.cell(row=i+2,column=12).value = str(related_script_file_name)

        
        ms_case_num_total = ms_case_num_total + ms_case_num
        auto_case_num_total = auto_case_num_total + auto_case_num
        same_case_num_total = same_case_num_total + same_case_num
        uncovernum_total = uncovernum_total + uncovernum

        if ms_case_num == 0:
            reason_1 = "查看MS用例的标签是否正确获取是否有对应的自动化脚本"
            reason.append(reason_1)

        if auto_case_num ==0:
            reason_2 = "所有的自动化用例的caseid都不正确或者没有对应的自动化脚本"
            reason.append(reason_2)
            
        if reason !=[]:
            sht.cell(row=i+2,column=12).value = str(reason)

    coverage_rate_total = f"{same_case_num_total/ms_case_num_total*100:.1f}%"
    sht.cell(row=i+3,column=1).value = "Total"
    sht.cell(row=i+3,column=2).value = coverage_rate_total
    sht.cell(row=i+3,column=3).value = uncovernum_total
    sht.cell(row=i+3,column=5).value = ms_case_num_total
    sht.cell(row=i+3,column=7).value = auto_case_num_total
    sht.cell(row=i+3,column=9).value = same_case_num_total
    return same_case_dic

class MS_Client(meterSphere_client):
    def get_testcase_id_from_ms():
        pass

class CollectCaseInfo(Upload):
    def get_caseid_and_tag(self,nodeIds, projectid):
        case_info = []
        res = self.client.get_testcases_from_nodeIds(nodeIds, projectid)
        # print(res)
        for i in range(len(res)) :
            # data = res[i].get('fields')
            # print(type(data))
            # print(f"-------------->{res[i].get('fields')[7]['value']}")
            try:
                if "Not Applicable" in str(res[i].get('fields')):
                    print(f"--------------------> 'case_id': {res[i].get('num')} Not need automatical <----------------------------------------")
                    pass
                else:
                    case_info.append({'id': res[i].get('id'), 'case_id': res[i].get('num'),'jama_id': res[i].get('customNum'), 'nodeid': res[i].get('nodeId'), 'functionModule': res[i].get('nodePath'), 'tags': res[i].get('tags')})
            except Exception as e:
                print(f"'case_id': {res[i].get('num')} Error: ------------->{str(e)}")
                pass
        return case_info

if __name__ == "__main__":
    current_file_path = os.path.realpath(__file__)
    pos = current_file_path.rfind(r'/')
    file_path_smoke = current_file_path[: pos + 1] + "collect_case_info_smoke.log"
    file_path_sanity = current_file_path[: pos + 1] + "collect_case_info_sanity.log"
    file_path_full = current_file_path[: pos + 1] + "collect_case_info_full.log"
    xls_path = current_file_path[: pos + 1] + "case_analysis_result_mcu.xlsx"
    xls_obj = openpyxl.load_workbook(xls_path)
    case_info = CollectCaseInfo()
    nodeIds = ["28c4aecb-b8b4-4df5-82f1-38a5bddefb40"]
    projectid = "1fbbeb47-6cc7-48bd-aafc-24349853e9d8" 
    mscase_info_list = case_info.get_caseid_and_tag([], projectid)
    # print(mscase_info_list)
    for i in range(len(mscase_info_list)):
        # print(mscase_info_list[i].get('fields'))
        print(r"case id:{},node id:{},function module:{},case tags:{}".format(mscase_info_list[i].get('case_id'),mscase_info_list[i].get('nodeid'),mscase_info_list[i].get('functionModule'),mscase_info_list[i].get('tags')))
    
    
    for tag in ["smoke","sanity","full"]:
        if tag == "smoke":
            file_path = file_path_smoke
        elif tag == "sanity":
            file_path = file_path_sanity
        elif tag == "full":
            file_path = file_path_full

        auto_case_info_dic = dealwith_collect_orignal_data(file_path)
        print(auto_case_info_dic)
        ms_case_info_dic = dealwith_ms_original_case(mscase_info_list,tag)
        print(ms_case_info_dic)
        result = compare_ms_and_auto_case_info(ms_case_info_dic,auto_case_info_dic,xls_obj,tag)
        print(result)
        sleep(2)

    xls_obj.save(xls_path)
    xls_obj.close()
    