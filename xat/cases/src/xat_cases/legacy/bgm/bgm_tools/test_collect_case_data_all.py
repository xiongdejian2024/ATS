import os
import sys
import re
from time import sleep
import pytest
import allure
import openpyxl

project_root = os.path.join(os.getcwd().split("sat")[0], "sat")
sys.path.append(project_root)

work_path_2 = os.path.join(
    os.getcwd().split("sat")[0], "sat", "ecu_simulator", "interface"
)
sys.path.append(work_path_2)

sys.path.append(os.getcwd())
sys.path.append(os.path.join(os.getcwd(), ".."))
sys.path.append(os.path.join(os.getcwd(), "../.."))
sys.path.append(os.path.join(os.getcwd(), "../../.."))

from xat_ecu.legacy.interface.ms.ms_lib import *
from xat_ecu.legacy.interface.ms.ms_to_jama import *

script_and_function_module_mapping = {
    "/BGM车云用例/远程诊断功能": [],
    "/BGM车云用例/数字钥匙功能": [],
    "/BGM车云用例/数据埋点": [],
    "/BGM车云用例/RVS功能/蓝牙数据上报": [],
    "/BGM车云用例/RVS功能/基础数据上报": [],
    "/BGM车云用例/GB32960功能": [],
    "/BGM车云用例/FOD功能": [],
    "/BGM车云用例/EDR功能": [],
    "/BGM车控车设/座椅功能": [],
    "/BGM车控车设/主动尾翼": [],
    "/BGM车控车设/主动尾翼/失效场景": [],
    "/BGM车控车设/整车模式/驻车舒享": [],
    "/BGM车控车设/整车模式/展车模式": [],
    "/BGM车控车设/整车模式/洗车模式": [],
    "/BGM车控车设/整车模式/维修模式": [],
    "/BGM车控车设/整车模式/使用模式": [],
    "/BGM车控车设/整车模式/车辆模式": [],
    "/BGM车控车设/雨刮功能": [],
    "/BGM车控车设/行人安全": [],
    "/BGM车控车设/尾门功能": [],
    "/BGM车控车设/外灯功能": [],
    "/BGM车控车设/胎压功能": [],
    "/BGM车控车设/锁模块功能": [],
    "/BGM车控车设/舒适进出": [],
    "/BGM车控车设/手套箱功能": [],
    "/BGM车控车设/内灯功能": [],
    "/BGM车控车设/空调控制": [],
    "/BGM车控车设/后视镜功能": [],
    "/BGM车控车设/防盗功能": [],
    "/BGM车控车设/方向盘功能": [],
    "/BGM车控车设/电动门功能": [],
    "/BGM车控车设/低压管理系统功能": [],
    "/BGM车控车设/灯光秀功能": [],
    "/BGM车控车设/充电口盖功能": [],
    "/BGM车控车设/乘客安全/安全带功能": [],
    "/BGM车控车设/车窗功能": [],
    "/BGM车控车设/WTI功能/座椅故障": [],
    "/BGM车控车设/WTI功能/主动尾翼故障": [],
    "/BGM车控车设/WTI功能/雨刮控制": [],
    "/BGM车控车设/WTI功能/钥匙": [],
    "/BGM车控车设/WTI功能/尾门报警提醒": [],
    "/BGM车控车设/WTI功能/外灯故障": [],
    "/BGM车控车设/WTI功能/门锁告警/电动门相关告警": [],
    "/BGM车控车设/WTI功能/空调故障": [],
    "/BGM车控车设/WTI功能/后视镜故障": [],
    "/BGM车控车设/WTI功能/防盗报警": [],
    "/BGM车控车设/WTI功能/方向盘故障": [],
    "/BGM车控车设/WTI功能/低压高压故障": [],
    "/BGM车控车设/WTI功能/充电口盖": [],
    "/BGM车控车设/WTI功能/乘客安全报警": [],
    "/BGM车控车设/WTI功能/车窗故障": [],
    "/BGM车控车设/WTI功能/TPMS_and_VMM": [],
    "/BGM车控车设/RelayControl功能": [],
    "/BGM_MCU/MCU性能/业务性能/车辆公告发出时长": [],
    "/BGM_MCU/MCU基础通讯/以太网通讯": [],
    "/BGM_MCU/MCU基础通讯/DoIP通讯/物理层": [],
    "/BGM_MCU/MCU基础功能/外灯": [],
    "/BGM_MCU/MCU基础功能/VMM": [],
    "/BGM_MCU/MCU 基础平台/NVM存储": [],
    "/BGM_MCU/MCU 基础平台/MCU诊断路由": [],
    "/BGM_MCU/MCU 基础平台/CDD打包功能/CDD性能用例": [],
    "/BGM FOTA用例/域控制器FOTA状态机": [],
    "/BGM FOTA用例/下载阶段Error_Code": [],
    "/BGM FOTA用例/退出FOTA流程": [],
    "/BGM FOTA用例/升级阶段Error_Code": [],
    "/BGM FOTA用例/前期条件/任务获取": [],
    "/BGM FOTA用例/前期条件/Error_Code重试": [],
    "/BGM FOTA用例/前期条件/Carmode": [],
    "/BGM FOTA用例/厂内OTA自检任务": [],
    "/BGM FOTA用例/厂内OTA条件检测": [],
    "/BGM FOTA用例/厂内OTA升级中": [],
    "/BGM FOTA用例/厂内OTA升级前": [],
    "/BGM FOTA用例/厂内OTA升级结束": [],
    "/BGM FOTA用例/厂内OTA任务触发": [],
    "/BGM FOTA用例/厂内OTA端到端/厂内OTA端到端": [],
    "/BGM FOTA用例/常规OTA端到端/常规OTA端到端": [],
    "/BGM FOTA用例/版本收集与上传": [],
    "/BGM FOTA用例/FOTA下载": [],
    "/BGM FOTA用例/FOTA升级中": [],
    "/BGM FOTA用例/FOTA升级前": [],
    "/BGM FOTA用例/FOTA取消": [],
    "/BGM BaseTech/诊断功能/诊断路由": [],
    "/BGM BaseTech/诊断功能/诊断服务": [],
    "/BGM BaseTech/诊断功能/诊断RoutineControl": [],
    "/BGM BaseTech/诊断功能/诊断IOControl": [],
    "/BGM BaseTech/诊断功能/诊断DID": [],
    "/BGM BaseTech/诊断功能/雷达标定": [],
    "/BGM BaseTech/诊断DTC": [],
    "/BGM BaseTech/信息安全/证书管理": [],
    "/BGM BaseTech/信息安全/系统安全": [],
    "/BGM BaseTech/信息安全/OBD防火墙": [],
    "/BGM BaseTech/通讯功能/以太网通讯": [],
    "/BGM BaseTech/通讯功能/LIN通讯": [],
    "/BGM BaseTech/软件平台/网络管理/PNC网络管理/PNC置位测试": [],
    "/BGM BaseTech/软件平台/网络管理/PNC网络管理/PNC路由测试": [],
    "/BGM BaseTech/软件平台/网络管理/LIN网络管理": [],
    "/BGM BaseTech/软件平台/数据存储": [],
    "/BGM BaseTech/软件平台/时间管理": [],
    "/BGM BaseTech/软件平台/日志管理": [],
    "/BGM BaseTech/软件平台/健康监控": [],
    "/BGM BaseTech/软件平台/电源管理": [],
    "/BGM BaseTech/工业化/离线刷写台": [],
    "/BGM BaseTech/车辆配置/CCP": [],
}


def dealwith_collect_orignal_data(file_path):
    pattern_0 = r"<Module\s(.*?)>"
    pattern_1 = r"_\d+"
    pattern_2 = r"\[(\d{5,})\]"
    poslist = []
    infodic = {}

    # r"C:\Users\hui.zhao\Desktop\Learning\Python\collect_case_info.log"
    with open(file_path, "r", encoding="utf-8") as f:
        file_cont = str(f.readlines())

    matches = re.findall(pattern_0, file_cont)
    for i in range(len(matches)):
        poslist.append(file_cont.find(matches[i]))
        if i > 0:
            if i < len(matches):
                str_1 = file_cont[poslist[i - 1] : poslist[i]]
            else:
                str_1 = file_cont[poslist[i - 1] : len(file_cont)]
            matches_1 = re.findall(pattern_1, str_1)
            matches_2 = re.findall(pattern_2, str_1)
            for n in range(len(matches_1)):
                matches_1[n] = int(matches_1[n].replace("_", ""))

            for m in range(len(matches_2)):
                matches_2[m] = int(matches_2[m].replace("case[", ""))

            pos = matches[i - 1].rfind("test_")
            infodic[matches[i - 1][pos:]] = matches_2 + matches_1
    return infodic


def clear_excel_content(sheet_obj):
    for l in range(2, 100):
        for r in range(1, 20):
            sheet_obj.cell(row=l, column=r).value = ""


def dealwith_ms_original_case(original_data_list, tag):
    global script_and_function_module_mapping
    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    ms_data_dic_1 = {}
    for i in range(len(original_data_list)):
        # if "1.4" in original_data_list[i].get("tags"):
        #     print("V1.4")
        #     continue

        for n in range(len(all_function_module_name_list)):
            if tag == "smoke":
                tag_list = json.loads(original_data_list[i].get("tags"))
                for num in tag_list:
                    for elem in ["smoke", "Smoke", "smoking", "Smoking"]:
                        if elem == num:
                            if all_function_module_name_list[n] in original_data_list[
                                i
                            ].get("functionModule"):
                                if (
                                    all_function_module_name_list[n]
                                    in ms_data_dic_1.keys()
                                ):
                                    if (
                                        original_data_list[i].get("case_id")
                                        not in ms_data_dic_1[
                                            all_function_module_name_list[n]
                                        ]
                                    ):
                                        ms_data_dic_1[
                                            all_function_module_name_list[n]
                                        ].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]] = [
                                        original_data_list[i].get("case_id")
                                    ]
            if tag == "checklist":
                tag_list = json.loads(original_data_list[i].get("tags"))
                for num in tag_list:
                    for elem in [
                        "smoke",
                        "Smoke",
                        "smoking",
                        "Smoking",
                        "Checklist",
                        "checklist",
                        "sanity",
                        "Sanity",
                    ]:
                        if elem == num:
                            if all_function_module_name_list[n] in original_data_list[
                                i
                            ].get("functionModule"):
                                if (
                                    all_function_module_name_list[n]
                                    in ms_data_dic_1.keys()
                                ):
                                    if (
                                        original_data_list[i].get("case_id")
                                        not in ms_data_dic_1[
                                            all_function_module_name_list[n]
                                        ]
                                    ):
                                        ms_data_dic_1[
                                            all_function_module_name_list[n]
                                        ].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]] = [
                                        original_data_list[i].get("case_id")
                                    ]
                                break

            if tag == "full":
                if all_function_module_name_list[n] in original_data_list[i].get(
                    "functionModule"
                ):
                    if all_function_module_name_list[n] in ms_data_dic_1.keys():
                        if (
                            original_data_list[i].get("case_id")
                            not in ms_data_dic_1[all_function_module_name_list[n]]
                        ):
                            ms_data_dic_1[all_function_module_name_list[n]].append(
                                original_data_list[i].get("case_id")
                            )
                    else:
                        ms_data_dic_1[all_function_module_name_list[n]] = [
                            original_data_list[i].get("case_id")
                        ]
    return ms_data_dic_1


def compare_ms_and_auto_case_info(ms_info, auto_script_info, excel_obj, tag):
    global cript_and_function_module_mapping

    if tag == "smoke":
        sht = excel_obj.get_sheet_by_name("smoke")
    elif tag == "checklist":
        sht = excel_obj.get_sheet_by_name("checklist")
    elif tag == "full":
        sht = excel_obj.get_sheet_by_name("full")

    clear_excel_content(sht)

    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    same_case_dic = {}
    coverage_rate_total = 0
    ms_case_num_total = 0
    auto_case_num_total = 0
    same_case_num_total = 0

    for i in range(len(all_function_module_name_list)):
        reason = []
        print(all_function_module_name_list[i])
        sht.cell(row=i + 2, column=1).value = all_function_module_name_list[i]
        if all_function_module_name_list[i] not in list(ms_info.keys()):
            ms_case = []
        else:
            ms_case = ms_info[all_function_module_name_list[i]]

        ms_case_num = len(ms_case)

        related_script_file_name = script_and_function_module_mapping[
            all_function_module_name_list[i]
        ]
        print(related_script_file_name)
        if type(related_script_file_name) == list:
            auto_case = []
            for elem in related_script_file_name:
                if elem in list(auto_script_info.keys()):
                    auto_case = auto_case + auto_script_info[elem]
        else:
            if related_script_file_name not in list(auto_script_info.keys()):
                auto_case = []
            else:
                auto_case = auto_script_info[related_script_file_name]
        auto_case_num = len(auto_case)

        same_case_list = []
        for ms_case_id in ms_case:
            for auto_case_id in auto_case:
                if ms_case_id == auto_case_id:
                    same_case_list.append(ms_case_id)

        same_case_num = len(same_case_list)

        if ms_case_num != 0:
            coverage_rate = len(same_case_list) / ms_case_num
            percentage = f"{coverage_rate*100:.1f}%"
        else:
            percentage = f"0%,可能此模块没有{tag}用例,也可能标签不正确，请进一步确认"

        sht.cell(row=i + 2, column=2).value = str(percentage)
        sht.cell(row=i + 2, column=3).value = str(ms_case_num)
        sht.cell(row=i + 2, column=4).value = str(ms_case)
        sht.cell(row=i + 2, column=5).value = str(auto_case_num)
        sht.cell(row=i + 2, column=6).value = str(auto_case)
        sht.cell(row=i + 2, column=7).value = str(len(same_case_list))
        sht.cell(row=i + 2, column=8).value = str(same_case_list)

        ms_not_include = [x for x in ms_case if x not in same_case_list]
        sht.cell(row=i + 2, column=9).value = str(ms_not_include)

        auto_not_include = [x for x in auto_case if x not in same_case_list]
        sht.cell(row=i + 2, column=10).value = str(auto_not_include)
        same_case_dic[all_function_module_name_list[i]] = same_case_list

        sht.cell(row=i + 2, column=11).value = str(related_script_file_name)

        ms_case_num_total = ms_case_num_total + ms_case_num
        auto_case_num_total = auto_case_num_total + auto_case_num
        same_case_num_total = same_case_num_total + same_case_num

        if ms_case_num == 0:
            reason_1 = "查看MS用例的标签是否正确获取是否有对应的自动化脚本"
            reason.append(reason_1)

        if auto_case_num == 0:
            reason_2 = "所有的自动化用例的caseid都不正确或者没有对应的自动化脚本"
            reason.append(reason_2)

        if reason != []:
            sht.cell(row=i + 2, column=12).value = str(reason)

    coverage_rate_total = f"{same_case_num_total/ms_case_num_total*100:.1f}%"
    sht.cell(row=i + 3, column=1).value = "Total"
    sht.cell(row=i + 3, column=2).value = coverage_rate_total
    sht.cell(row=i + 3, column=3).value = ms_case_num_total
    sht.cell(row=i + 3, column=5).value = auto_case_num_total
    sht.cell(row=i + 3, column=7).value = same_case_num_total
    return same_case_dic


class MS_Client(meterSphere_client):
    def get_testcase_id_from_ms():
        pass


class CollectCaseInfo(Upload):
    def get_caseid_and_tag(self, nodeIds, projectid):
        case_info = []
        res = self.client.get_testcases_from_nodeIds(nodeIds, projectid)
        # print(res)
        for i in range(len(res)):
            # data = res[i].get('fields')
            # print(type(data))
            # print(f"-------------->{res[i].get('fields')[7]['value']}")
            try:
                if "Not Applicable" in str(res[i].get("fields")):
                    print(
                        f"--------------------> 'case_id': {res[i].get('num')} Not need automatical <----------------------------------------"
                    )
                    pass
                else:
                    case_info.append(
                        {
                            "id": res[i].get("id"),
                            "case_id": res[i].get("num"),
                            "jama_id": res[i].get("customNum"),
                            "nodeid": res[i].get("nodeId"),
                            "functionModule": res[i].get("nodePath"),
                            "tags": res[i].get("tags"),
                        }
                    )
            except Exception as e:
                print(f"'case_id': {res[i].get('num')} Error: ------------->{str(e)}")
                pass
        return case_info


if __name__ == "__main__":
    current_file_path = os.path.realpath(__file__)
    pos = current_file_path.rfind(r"/")
    file_path_smoke = current_file_path[: pos + 1] + "collect_case_info_smoke.log"
    file_path_checklist = (
        current_file_path[: pos + 1] + "collect_case_info_checklist.log"
    )
    file_path_full = current_file_path[: pos + 1] + "collect_case_info_full.log"
    xls_path = current_file_path[: pos + 1] + "case_analysis_result.xlsx"
    xls_obj = openpyxl.load_workbook(xls_path)
    case_info = CollectCaseInfo()
    nodeIds = ["28c4aecb-b8b4-4df5-82f1-38a5bddefb40"]
    projectid = "1fbbeb47-6cc7-48bd-aafc-24349853e9d8"
    mscase_info_list = case_info.get_caseid_and_tag([], projectid)
    # print(mscase_info_list)
    for i in range(len(mscase_info_list)):
        # print(mscase_info_list[i].get('fields'))
        print(
            r"case id:{},node id:{},function module:{},case tags:{}".format(
                mscase_info_list[i].get("case_id"),
                mscase_info_list[i].get("nodeid"),
                mscase_info_list[i].get("functionModule"),
                mscase_info_list[i].get("tags"),
            )
        )

    for tag in ["smoke", "checklist", "full"]:
        if tag == "smoke":
            file_path = file_path_smoke
        elif tag == "checklist":
            file_path = file_path_checklist
        elif tag == "full":
            file_path = file_path_full

        auto_case_info_dic = dealwith_collect_orignal_data(file_path)
        print(auto_case_info_dic)
        ms_case_info_dic = dealwith_ms_original_case(mscase_info_list, tag)
        print(ms_case_info_dic)
        result = compare_ms_and_auto_case_info(
            ms_case_info_dic, auto_case_info_dic, xls_obj, tag
        )
        print(result)
        sleep(2)

    xls_obj.save(xls_path)
    xls_obj.close()
