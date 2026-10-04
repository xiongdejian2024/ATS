import os
import sys
import re
from time import sleep
import pytest
import allure
import openpyxl

project_root = os.path.join(os.getcwd().split("sat")[0], 'sat')
sys.path.append(project_root)

work_path_2 = os.path.join(os.getcwd().split("sat")[0], 'sat', 'ecu_simulator', 'interface')
sys.path.append(work_path_2)

sys.path.append(os.getcwd())
sys.path.append(os.path.join(os.getcwd(), ".."))
sys.path.append(os.path.join(os.getcwd(), "../.."))
sys.path.append(os.path.join(os.getcwd(), "../../.."))

from xat_ecu.legacy.interface.ms.ms_lib import *
from xat_ecu.legacy.interface.ms.ms_to_jama import *

script_and_function_module_mapping = {
    "/BGM车云用例/RVS功能/蓝牙数据上报/车辆模式": ["test_ble_vehicle_mode.py"],
    "/BGM车云用例/RVS功能/蓝牙数据上报/车身状态": ["test_ble_vehicle_body.py"],
    "/BGM车云用例/RVS功能/蓝牙数据上报/充电状态": ["test_ble_eic_charging.py"],
    "/BGM车云用例/RVS功能/蓝牙数据上报/车辆状态": ["test_ble_vehicle_status.py"],
    "/BGM车云用例/RVS功能/蓝牙数据上报/AVP": ["test_ble_avp.py"],
    "/BGM车云用例/RVS功能/基础数据上报/ECU版本收集":[],
    "/BGM车云用例/RVS功能/基础数据上报/车辆模式":["test_vehicle_mode.py"],
    "/BGM车云用例/RVS功能/基础数据上报/车身状态":["test_vehicle_body.py"],
    "/BGM车云用例/RVS功能/基础数据上报/空调状态":["test_cabin_status.py"],
    "/BGM车云用例/RVS功能/基础数据上报/GNSS&仪表显示":["test_gnss_travel.py"],
    "/BGM车云用例/RVS功能/基础数据上报/行驶状态":["test_driving_status.py"],
    "/BGM车云用例/RVS功能/基础数据上报/充电状态":["test_eic_charging.py"],
    "/BGM车云用例/RVS功能/基础数据上报/车辆状态":["test_vehicle_status.py"],
    "/BGM车云用例/RVS功能/基础数据上报/数字钥匙":[],
    "/BGM车云用例/RVS功能/基础数据上报/EOL数据":[],
    "/BGM车云用例/RVS功能/基础数据上报/灯光状态":["test_light.py"],
    "/BGM车云用例/RVS功能/基础数据上报/WTI状态":[],
    "/BGM车云用例/数字钥匙功能/PEPS/近车自动解锁":['test_peps_approach.py', 'test_peps_approach_abc.py'],
    "/BGM车云用例/数字钥匙功能/PEPS/离车侧门尾门自动关闭落锁":['test_peps_walk_away.py'],
    "/BGM车云用例/数字钥匙功能/PEPS/触碰控制":['test_peps_PE.py'],
    "/BGM车云用例/数字钥匙功能/PEPS/启动授权":['test_peps_start_up_authorization.py'],
    "/BGM车云用例/数字钥匙功能/RKE/蓝牙解闭锁":['test_rke_lock_control.py', 'test_rke_lock_control_abc.py'],
    "/BGM车云用例/数字钥匙功能/RKE/蓝牙车窗控制":['test_rke_window_control.py'],
    "/BGM车云用例/数字钥匙功能/RKE/MaxSOC控制":["test_max_soc_abc.py"],
    "/BGM车云用例/数字钥匙功能/RKE/尾门控制":['test_rke_tailgate_control.py', 'test_rke_tailgate_control_abc.py'],
    "/BGM车云用例/数字钥匙功能/RKE/APA五门全关":['test_rke_apa_close_door.py'],
    "/BGM车云用例/数字钥匙功能/RKE/寻车控制":['test_rke_car_locator.py', 'test_rke_car_locator_abc.py'],
    "/BGM车云用例/数字钥匙功能/RKE/远程更新BNCM个人化数据":["test_rke_se_abc.py"],
    "/BGM车云用例/数字钥匙功能/RKE/充电桩蓝牙充电口盖控制":['test_charge_lid_abc.py'],
    "/BGM车云用例/数字钥匙功能/RKE/APA控制":['test_apa_control.py'],
    "/BGM车云用例/数字钥匙功能/DK网络管理":['test_digital_key_NM.py'], 
    "/BGM车云用例/数字钥匙功能/账号控制与同步":['test_key_id_sync_abc.py', 'test_white_list_control.py', 'test_white_list_control_abc.py','test_nfc.py','test_nfc_learning_abc.py'],
    "/BGM车云用例/数字钥匙功能/UWB":['test_peps_PE_abc.py'], 
    "/BGM车云用例/远程诊断功能":['test_remote_diag.py'],
    "/BGM车云用例/FOD功能/端到端用例":["test_fod_bench_e2e.py",'test_network_sleep_fod_vehicle_e2e.py'],
    "/BGM车云用例/FOD功能/CCP Master状态机":["test_ccp_status.py"],
    "/BGM车云用例/FOD功能/唤醒和条件检测":["test_ccp_wake_and_condition_check.py"],
    "/BGM车云用例/FOD功能/激活过程_下切VMM至Convience和进FOTA":["test_set_vmm_conn_enter_fotamode.py"],
    "/BGM车云用例/FOD功能/激活过程_写CCP和切VMM至active":["test_write_ccp_set_vmm_active.py"],
    "/BGM车云用例/FOD功能/激活过程_等待激活":["test_wait_active.py"],
    "/BGM车云用例/FOD功能/激活过程_结果确认":[],
    "/BGM车云用例/FOD功能/激活过程_CCP回滚":["test_ccp_rollback.py"],
    "/BGM车云用例/FOD功能/激活过程_退FOTAMODE":["test_ccp_quit_fotamode.py"],
    "/BGM车云用例/FOD功能/激活过程_退出流程":["test_ccp_exit_process.py"],
    "/BGM车云用例/远程控制":["test_car_trace.py","test_charge_control.py","test_climate.py","test_lock.py","test_max_climate.py","test_passtest_lock.py","test_remote_authorization.py","test_seat.py","test_steer_wheel_heating.py","test_tailgate.py","test_window.py"],
    "/BGM车云用例/数据埋点":[]
}


def dealwith_collect_orignal_data(file_path):
    pattern_0 = r"<Module\s(.*?)>"
    pattern_1 = r"_\d+"
    pattern_2 = r"\[(\d{5,})\]"
    poslist = []
    infodic = {}
    
    # r"C:\Users\hui.zhao\Desktop\Learning\Python\collect_case_info.log"
    with open(file_path,'r', encoding='utf-8') as f:  
        file_cont = str(f.readlines())   

    matches = re.findall(pattern_0, file_cont)
    for i in range(len(matches)):
        poslist.append(file_cont.find(matches[i]))
        if i >0 :
            if i < len(matches):
                str_1 = file_cont[poslist[i-1]:poslist[i]]
            else:
                str_1= file_cont[poslist[i-1]:len(file_cont)]
            matches_1 = re.findall(pattern_1, str_1)
            matches_2 = re.findall(pattern_2, str_1)
            for n in range(len(matches_1)):
                matches_1[n] = int(matches_1[n].replace("_",""))

            for m in range(len(matches_2)):
                matches_2[m] = int(matches_2[m].replace("case[",""))

            pos = matches[i-1].rfind("test_")
            infodic[matches[i-1][pos:]] = matches_2 + matches_1
    return infodic


def clear_excel_content(sheet_obj):
    for l in range(2,100):
        for r in range(1,20):
            sheet_obj.cell(row=l,column=r).value = ""

    
def dealwith_ms_original_case(original_data_list,tag):
    global script_and_function_module_mapping
    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    ms_data_dic_1 = {}
    for i in range(len(original_data_list)):
        if "1.4_Only" in original_data_list[i].get("tags"):
            continue
        for n in range(len(all_function_module_name_list)):
            priority_list = original_data_list[i].get("priority") # json.loads(original_data_list[i].get("priority"))
            if tag == "smoke":
                if priority_list == "P0":
                    if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                                if all_function_module_name_list[n] in ms_data_dic_1.keys():
                                    if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                                        ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]
                                break
            if tag == "sanity":
                if priority_list == "P0" or priority_list == "P1":
                    if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                                if all_function_module_name_list[n] in ms_data_dic_1.keys():
                                    if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                                        ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                                else:
                                    ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]
                                break
            if tag == "full":
                if all_function_module_name_list[n] in original_data_list[i].get("functionModule"):
                    if all_function_module_name_list[n] in ms_data_dic_1.keys():
                        if original_data_list[i].get("case_id") not in ms_data_dic_1[all_function_module_name_list[n]]:
                            ms_data_dic_1[all_function_module_name_list[n]].append(original_data_list[i].get("case_id"))
                    else:
                        ms_data_dic_1[all_function_module_name_list[n]]= [original_data_list[i].get("case_id")]
    return ms_data_dic_1
    


def compare_ms_and_auto_case_info(ms_info,auto_script_info,excel_obj,tag):
    global cript_and_function_module_mapping
    
    if tag == "smoke":
        sht=excel_obj.get_sheet_by_name("smoke")
    elif tag == "sanity":
        sht=excel_obj.get_sheet_by_name("sanity")
    elif tag == "full":
        sht=excel_obj.get_sheet_by_name("full")
    
    clear_excel_content(sht)
        
    all_function_module_name_list = list(script_and_function_module_mapping.keys())
    same_case_dic = {}
    coverage_rate_total = 0
    ms_case_num_total = 0
    auto_case_num_total = 0
    same_case_num_total = 0
    uncovernum_total = 0
 
    for i in range(len(all_function_module_name_list)):
        reason = []
        print(all_function_module_name_list[i])
        sht.cell(row=i+2,column=1).value = all_function_module_name_list[i]
        if all_function_module_name_list[i] not in list(ms_info.keys()):
            ms_case = []
        else:
            ms_case = ms_info[all_function_module_name_list[i]]

        ms_case_num = len(ms_case)
            
        related_script_file_name = script_and_function_module_mapping[all_function_module_name_list[i]]
        print(related_script_file_name)
        if type(related_script_file_name) == list:
            auto_case = []
            for elem in  related_script_file_name:
                if elem in list(auto_script_info.keys()):
                    auto_case = auto_case + auto_script_info[elem]
        else:
            if related_script_file_name not in list(auto_script_info.keys()):
                auto_case = []
            else:
                auto_case = auto_script_info[related_script_file_name]
        auto_case_num = len(auto_case)
        print(f"---------------------------------->auto_case id list{auto_case}")

        same_case_list = []
        for ms_case_id in ms_case:
            for auto_case_id in auto_case:
                if ms_case_id == auto_case_id and ms_case_id not in same_case_list:
                    same_case_list.append(ms_case_id)
        
        same_case_num = len(same_case_list)
        print(f"---------------------------------->same_case_num id list{same_case_num}")
        
        if ms_case_num != 0:
            coverage_rate = len(same_case_list)/ms_case_num    
            percentage = f"{coverage_rate*100:.1f}%"
        else:
            percentage = f"0%,可能此模块没有{tag}用例,也可能标签不正确，请进一步确认"
        
        print(f"---------------------------------->percentage:{percentage}")
        
        
        sht.cell(row=i+2,column=2).value = str(percentage)
        auto_not_include = [x for x in ms_case if (x not in same_case_list)]
        uncovernum = len(auto_not_include)

        sht.cell(row=i+2,column=3).value = str(uncovernum)
        sht.cell(row=i+2,column=4).value = str(auto_not_include)
        sht.cell(row=i+2,column=5).value = str(ms_case_num)
        sht.cell(row=i+2,column=6).value = str(ms_case)
        sht.cell(row=i+2,column=7).value = str(auto_case_num)
        sht.cell(row=i+2,column=8).value = str(auto_case)
        sht.cell(row=i+2,column=9).value = str(len(same_case_list))
        sht.cell(row=i+2,column=10).value = str(same_case_list)
        
        ms_not_include = [x for x in auto_case if x not in same_case_list]
        sht.cell(row=i+2,column=11).value = str(ms_not_include)
        
        same_case_dic[all_function_module_name_list[i]] = same_case_list   
        sht.cell(row=i+2,column=12).value = str(related_script_file_name)

        
        ms_case_num_total = ms_case_num_total + ms_case_num
        auto_case_num_total = auto_case_num_total + auto_case_num
        same_case_num_total = same_case_num_total + same_case_num
        uncovernum_total = uncovernum_total + uncovernum

        if ms_case_num == 0:
            reason_1 = "查看MS用例的标签是否正确获取是否有对应的自动化脚本"
            reason.append(reason_1)

        if auto_case_num ==0:
            reason_2 = "所有的自动化用例的caseid都不正确或者没有对应的自动化脚本"
            reason.append(reason_2)
            
        if reason !=[]:
            sht.cell(row=i+2,column=12).value = str(reason)

    if ms_case_num_total !=0:
        coverage_rate_total = f"{same_case_num_total/ms_case_num_total*100:.1f}%"
    sht.cell(row=i+3,column=1).value = "Total"
    sht.cell(row=i+3,column=2).value = coverage_rate_total
    sht.cell(row=i+3,column=3).value = uncovernum_total
    sht.cell(row=i+3,column=5).value = ms_case_num_total
    sht.cell(row=i+3,column=7).value = auto_case_num_total
    sht.cell(row=i+3,column=9).value = same_case_num_total
    return same_case_dic


    
class MS_Client(meterSphere_client):
    def get_testcase_id_from_ms():
        pass

class CollectCaseInfo(Upload):
    def get_caseid_and_tag(self,nodeIds, projectid):
        case_info = []
        res = self.client.get_testcases_from_nodeIds(nodeIds, projectid)
        for i in range(len(res)) :
            try:
                if "Not Applicable" in str(res[i].get('fields')):
                    print(f"--------------------> 'case_id': {res[i].get('num')} Not need automatical <----------------------------------------")
                    pass
                else:
                    case_info.append({'id': res[i].get('id'), 'case_id': res[i].get('num'),'jama_id': res[i].get('customNum'), 'nodeid': res[i].get('nodeId'), 'functionModule': res[i].get('nodePath'), 'priority':res[i].get('priority'),'tags': res[i].get('tags')})
            except Exception as e:
                print(f"'case_id': {res[i].get('num')} Error: ------------->{str(e)}")
                pass
        return case_info

if __name__ == "__main__":
    current_file_path = os.path.realpath(__file__)
    pos = current_file_path.rfind(r'/')
    file_path_smoke = current_file_path[: pos + 1] + "collect_case_info_smoke.log"
    file_path_sanity = current_file_path[: pos + 1] + "collect_case_info_sanity.log"
    file_path_full = current_file_path[: pos + 1] + "collect_case_info_full.log"
    xls_path = current_file_path[: pos + 1] + "case_analysis_result_cloud.xlsx"
    xls_obj = openpyxl.load_workbook(xls_path)
    case_info = CollectCaseInfo()
    nodeIds = ["28c4aecb-b8b4-4df5-82f1-38a5bddefb40"]
    projectid = "1fbbeb47-6cc7-48bd-aafc-24349853e9d8" 
    mscase_info_list = case_info.get_caseid_and_tag([], projectid)
    # print(mscase_info_list)
    for i in range(len(mscase_info_list)):
        # print(mscase_info_list[i].get('fields'))
        print(r"case id:{},node id:{},function module:{},case tags:{}".format(mscase_info_list[i].get('case_id'),mscase_info_list[i].get('nodeid'),mscase_info_list[i].get('functionModule'),mscase_info_list[i].get('tags')))
    
    
    for tag in ["smoke","sanity","full"]:
        if tag == "smoke":
            file_path = file_path_smoke
        elif tag == "sanity":
            file_path = file_path_sanity
        elif tag == "full":
            file_path = file_path_full

        auto_case_info_dic = dealwith_collect_orignal_data(file_path)
        # print(auto_case_info_dic)
        ms_case_info_dic = dealwith_ms_original_case(mscase_info_list,tag)
        # print(ms_case_info_dic)
        result = compare_ms_and_auto_case_info(ms_case_info_dic,auto_case_info_dic,xls_obj,tag)
        # print(result)
        sleep(2)

    xls_obj.save(xls_path)
    xls_obj.close()
    
