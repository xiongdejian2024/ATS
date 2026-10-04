
import os
import sys

project_root = os.path.join(os.getcwd().split("sat")[0], 'sat')
sys.path.append(project_root)
from xat_ecu.legacy.driver.adb_client import Adb
from xat_ecu.legacy.common.logger import logger  as logging
from xat_cases.legacy.soa.case_helper.codesrc.components.pressure.utils import get_script_path, make_subprocess



import docx

def chech_event(base_dir):
    try:
        files = [os.path.join(base_dir, file) for file in os.listdir(base_dir) if
            os.path.join(base_dir, file).endswith('docx')]
        files.sort()
        path = files[-1]
        string = []
        error_flag = False
        doc = docx.Document(path)  # Creating word reader object.
        nrows = len(doc.tables[0].rows)
        ncols = len(doc.tables[0].columns)
        for p in doc.paragraphs:
            name = p.style.name
            if name.startswith('Heading'):
                if p.text.startswith("Method/Event的非正常调用check"):
                    for i in range(1, nrows):
                        for j in range(2, ncols):
                            string.append(doc.tables[0].cell(i, j).text.strip())
                    for z in string:
                        if int(z) != 0:
                            return error_flag, string
                    break
        return not error_flag, string
    
    except Exception as e:
        logging.error(str(e))


def read_doc(base_dir):
    try:

        files = [os.path.join(base_dir, file) for file in os.listdir(base_dir) if
                 os.path.join(base_dir, file).endswith('docx')]
        files.sort()
        path = files[-1]
        # 过滤到 异常的内容
        error_list = ['未过滤到initialize', "coredump", "tombstones"]
        string = ""
        error_flag = False
        doc = docx.Document(path)  # Creating word reader object.
        for para in doc.paragraphs:
            line = para.text
            lis = [True for i in error_list if i in line]
            if len(lis):
                error_flag = True
            string += line + "\n"
            if str(line).startswith("总成功率（去除成功个数"):
                break
        return not error_flag, string

    except Exception as e:
        logging.error(str(e))
        return False, str(e)


def get_paragraph_name(docx_file, table_index):
    """
    获取表格所在段落的标题名
    """
    doc = docx.Document(docx_file).paragraphs
    for p in range(len(doc)):
        name = doc[p].style.name
        if name.startswith('Heading'):
            return doc[p+table_index].text
        else:
            p += 1


def check_timeout_process(base_dir):
    """
    校验四域进程连接的最大值时间
    """
    event_Warning = ""
    event_Error = ""
    error_flag = False
    try:
        files = [os.path.join(base_dir, file) for file in os.listdir(base_dir) if
            os.path.join(base_dir, file).endswith('docx')]
        files.sort()
        path = files[-1]
        doc = docx.Document(path)

        for i in range(len(doc.tables)):
            paragraph_name = get_paragraph_name(path, i)
            nrows = len(doc.tables[i].rows)
            ncols = len(doc.tables[i].columns)
            for row in range(1, nrows):
                if doc.tables[0].cell(0, 0).text.strip() == "Proxy所在进程":
                    if paragraph_name == 'tcam':
                        if float(doc.tables[i].cell(row, 9).text.strip()) > 2*4000.0:
                            err = f"tcam端进程{doc.tables[i].cell(row, 0).text.strip()}连接目标服务{doc.tables[i].cell(row, 1).text.strip()}的最大时间{doc.tables[i].cell(row, 9).text.strip()}，超过8s"
                            event_Error += err + '\n'
                            logging.error(err)
                    else:
                        if 2000.0 <= float(doc.tables[i].cell(row, 9).text.strip()) <= 4000.0:
                            warn = f"进程{doc.tables[i].cell(row, 0).text.strip()}连接目标服务{doc.tables[i].cell(row, 1).text.strip()}的最大时间{doc.tables[i].cell(row, 9).text.strip()}，在2-4s之间"
                            event_Warning += warn + '\n'
                            logging.warning(warn)
                        elif float(doc.tables[i].cell(row, 9).text.strip()) > 4000.0:
                            err = f"进程{doc.tables[i].cell(row, 0).text.strip()}连接目标服务{doc.tables[i].cell(row, 1).text.strip()}的最大时间{doc.tables[i].cell(row, 9).text.strip()}，超过4s"
                            event_Error += err + '\n'
                            logging.error(err)

        if event_Error:
            return error_flag,event_Error
        else:
            return not error_flag,event_Warning
    except Exception as e:
        logging.error(e)
    

def run_as_root(cmd, pwd='123'):
    import pexpect
    child = pexpect.spawn(f'sudo {cmd}')
    child.sendline(pwd)
    child.expect('$')
    return child.before.decode()


def reset_log(log_path):
    if 'bgm' in log_path:
        os.system(f"chmod 777 {log_path}/*")
        if not os.path.isdir(os.path.join(log_path, 'bootes')):
            os.makedirs(os.path.join(log_path, 'bootes'))
        os.system(f"rm {os.path.join(log_path, 'bootes/*')}")
        files = os.popen(f"ls {os.path.join(log_path, 'jetlog_bts*')}").readlines()
        for file in files:
            if '.zst' in file:
                run_as_root(f"mv {file} {os.path.join(log_path, 'bootes')}")
    elif 'tcam' in log_path:
        if not os.path.isdir(os.path.join(log_path, 'log', 'bootes')):
            os.makedirs(os.path.join(log_path, 'log', 'bootes'))
        os.system(f"rm {os.path.join(log_path, 'log', 'bootes/*')}")
        files = os.popen(f"ls {os.path.join(log_path, 'log', 'jetlog_bts*')}").readlines()
        for file in files:
            if '.zst' in file:
                run_as_root(f"mv {os.path.join(log_path, 'log', file)} {os.path.join(log_path, 'log', 'bootes')}")
    elif 'cdc' in log_path:
        if not os.path.isdir(os.path.join(log_path, 'bootes')):
            os.makedirs(os.path.join(log_path, 'bootes'))
        run_as_root(f"mv {os.path.join(log_path, 'bootes/*')} {os.path.join(log_path, 'log/*')}")
        files = os.popen(f"ls {os.path.join(log_path, 'log')}").readlines()
        for file in files:
            if '.gz' in file:
                run_as_root(f"mv {os.path.join(log_path, 'log', file)} {os.path.join(log_path, 'bootes')}")
            if 'default' in file:
                run_as_root(f"mv {os.path.join(log_path, 'log', file)} {os.path.join(log_path, 'bootes')}")
    elif 'acu' in log_path:
        run_as_root(f"mv {os.path.join(log_path, 'acu_slave_bootes/*')} {os.path.join(log_path, 'bootes/*')}")


def get_acu_slave_log(user, bgm_addr, bgm_pwd, log_path, devices='99c4cd39'):
    script_path = os.path.join(project_root, 'soa_lib/codesrc/components/pressure/domain')
    cmds = [
                f'sshpass -p {bgm_pwd} scp -r {script_path}/script/port_mapping/config_acu_slave.sh  {user}@{bgm_addr}:/tmp/jiduer/',
                f'sshpass -p {bgm_pwd} ssh {user}@{bgm_addr} "chmod 755 /tmp/jiduer/config_acu_slave.sh"',
                f'sshpass -p {bgm_pwd} ssh {user}@{bgm_addr} "/tmp/jiduer/config_acu_slave.sh"'
            ]

    try:
        logging.info("清空原有acu从板日志")
        files = Adb(devices).shell("rm /mnt/sdcard/acu_slave_bootes/*")
        logging.info("下载acu从板日志")
        for cmd in cmds:
            make_subprocess(cmd)
        files = Adb(devices).shell("ls /mnt/sdcard/acu_slave_bootes")
        if not os.path.isdir(os.path.join(log_path, "acu_slave_bootes")):
            os.mkdir(os.path.join(log_path, "acu_slave_bootes"))
        for i in files.split(".log"):
            if "bootes" in i:
                os.system(f"adb -s {devices} pull /mnt/sdcard/acu_slave_bootes/{i.split('0m')[-1]+'.log'} {log_path}/acu_slave_bootes/")

    except Exception as e:
            print(f'acu从板日志下载失败{e}')


# get_acu_slave_log("jiduer", "169.254.19.1", "bgm@Axzfr778", "/root/TestDev/pressure_log/240322-100751-1-bootes1.4.1r703-172.18.128.138/1/acu")