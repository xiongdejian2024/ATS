# -*- coding: utf-8 -*-

"""
@Time    : 2024/06/03 11:55
@Author  : lei.tao
@Email   : lei.tao@jiduauto.com
"""
import requests
import json
import os
import time
import xlrd2
import openpyxl
import pandas as pd
from xat_ecu.legacy.common.logger import logger, Logger
from xat_ecu.legacy.interface.feishu.feishu_api import feishu_api

current_path = os.path.dirname(os.path.realpath(__file__))


class TestSignalCheck(feishu_api):
    def __super__(self) -> None:
        super(feishu_api())
        os.system("""ps -ef|grep -i diagd|awk '{printf $2"\n"}'|xargs kill -9;ps -ef|grep -i diagd""")
        self.get_user_access_token()

    def get_file_info(self,name,path="open-apis/sheets/v3/spreadsheets"):
        """
        获取文件的数据表信息
        """
        headers = {'Content-Type': 'application/json; charset=utf-8',
                   "Authorization": f"Bearer {self.tenant_access_token}"}
        url = self.server_name + path + f":{name}/sheets/query"

        response = requests.get(url, headers=headers, timeout=5)
        resp_body = json.loads(response.text)
        print(resp_body)
        print(self.tenant_access_token)


def get_caseinfo_from_sdbsignals(file, sheet_name):
    """
    获取sdbsignals表格中的服务接口信号信息
    """
    excel_path = current_path.split("sat")[0] + file
    wb1 = xlrd2.open_workbook(excel_path, encoding_override="utf-8")
    sheet = wb1.sheet_by_name(sheet_name)
    testcases_list = {}
    num_rows = sheet.nrows
    for i in range(1, num_rows):
        # GB32960Service服务的接口GB32960Data为周期上报，跟信号是否跳变没关系
        if sheet.cell_value(i, 0) not in ["ACUFaultInfoService", "GB32960Service"]:
            if not sheet.cell_value(i, 1).lower().startswith("get"):
                if not sheet.cell_value(i, 2).lower().endswith("chks") and not sheet.cell_value(i, 2).lower().endswith("cntr") and not sheet.cell_value(i, 2).lower().endswith("_ub"):
                    if sheet.cell_value(i, 0) == "ACCService":
                        testcases_begin = (sheet.cell_value(i, 0),sheet.cell_value(i, 1),sheet.cell_value(i, 3),sheet.cell_value(i, 4))

                        if testcases_begin not in testcases_list.keys():
                            testcases_list[testcases_begin] = [sheet.cell_value(i, 2)]
                        else:
                            testcases_list[testcases_begin].append(sheet.cell_value(i, 2))
                    else:
                        if not sheet.cell_value(i, 1).lower().startswith("set"):
                            testcases_begin = (sheet.cell_value(i, 0),sheet.cell_value(i, 1),sheet.cell_value(i, 3),sheet.cell_value(i, 4))

                            if testcases_begin not in testcases_list.keys():
                                testcases_list[testcases_begin] = [sheet.cell_value(i, 2)]
                            else:
                                testcases_list[testcases_begin].append(sheet.cell_value(i, 2))

    return testcases_list


def put_caseinfo_into_ms(sdbsignals_file,ms_file):
    """
    把sdbsignals表格中的服务接口信号信息转换成ms格式的用例
    """
    wb = openpyxl.load_workbook(current_path.split("sat")[0] + ms_file)
    sheet = wb.active
    header = []
    if not header:
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(1, c).value
            if val:
                header.append(val)
    ms_excel_path = current_path.split("sat")[0] + f"ms_{ms_file}"
    if os.path.exists(ms_excel_path):
        os.system(f"rm -rf {ms_excel_path}")
    wb.save(filename=ms_excel_path)
    # 只需要关注Mars1 sheet页
    for sdbsignals_sheet in ["Mars1_SDBSignals"]:
        data_list = get_caseinfo_from_sdbsignals(sdbsignals_file, sdbsignals_sheet)
        print(f"需要更新的{sdbsignals_sheet}信号数据：{data_list}")

        data = pd.read_excel(ms_excel_path, keep_default_na=True)
        data_dict = data.to_dict()
        print(data_dict)
        for key in enumerate(data_list.items()):
            print(f"开始插入数据：{key}")
            service_name = key[-1][0][0]
            interface_name = key[-1][0][1]
            signal_name = '，'.join(key[-1][-1])
            network_name = key[-1][0][2]
            message_name = key[-1][0][3]
            col = len(data)
            case_name = f'{service_name}::{interface_name}_{network_name}::{message_name}同帧非相关信号跳变无异常event'
            step_name = f'【1】取{network_name}::{message_name}中{signal_name}以外的信号赋值0和1，校验是否产生异常{interface_name}事件'
            if case_name in data_dict['用例名称'].values():
                for i in range(len(data_dict['ID'])):
                    if case_name == data_dict['用例名称'][i]:
                        if step_name == data_dict['步骤描述'][i]:
                            if sdbsignals_sheet.startswith("Venus_") and data_dict['适用车型'][i] == "[\"Mars One\"]":
                                print("有venus车型的信号完全一致，只修改适用车型字段")
                                data.loc[i, '适用车型'] = "[\"Mars One\", \"Venus\"]"
                            elif sdbsignals_sheet.startswith("Mars1_") and data_dict['适用车型'][i] == "[\"Venus\"]":
                                print("有mars1车型的信号完全一致，只修改适用车型字段")
                                data.loc[i, '适用车型'] = "[\"Mars One\", \"Venus\"]"
                            else:
                                print("车型的信号完全一致，不需要增加")

                        else:
                            if sdbsignals_sheet.startswith("Venus_") and data_dict['适用车型'][i] == "[\"Venus\"]":
                                print("有Venus车型的同帧信号有变更，修改该case")
                                data.loc[i, '步骤描述'] = step_name
                            elif sdbsignals_sheet.startswith("Mars1_") and data_dict['适用车型'][i] == "[\"Mars One\"]":
                                print("有mars1车型的同帧信号有变更，修改该case")
                                data.loc[i, '步骤描述'] = step_name
                            else:
                                print("有venus车型的同帧信号有变更，增加该case")
                                data.loc[key[0]+col, '用例名称'] = case_name
                                data.loc[key[0]+col, '步骤描述'] = step_name
                                data.loc[key[0]+col, '预期结果'] = f"【1】同帧非相关信号跳变无异常{interface_name}事件"
                                # 其他列填入默认值
                                data.loc[key[0]+col, '所属模块'] = "/SOA服务接口/通用场景/信号跳变校验event事件场景/AllService"
                                data.loc[key[0]+col, '整车版本'] = "[\"V2.0.0 Beta2\"]"
                                data.loc[key[0]+col, '编辑模式'] = "STEP"
                                data.loc[key[0]+col, '标签'] = "two_domain"
                                data.loc[key[0]+col, '用例类型'] = "接口用例"
                                data.loc[key[0]+col, '责任人'] = "陶磊 Lei(SOA)"
                                data.loc[key[0]+col, '用例等级'] = "P3"
                                data.loc[key[0]+col, '适用范围'] = "Full"
                                data.loc[key[0]+col, '业务级别'] = "[PlatformTest-SOA]"
                                data.loc[key[0]+col, 'Automation'] = "Automated"
                                data.loc[key[0]+col, '测试平台'] = "Bench"
                                data.loc[key[0]+col, "函数名称"] = interface_name
                                if sdbsignals_sheet.startswith("Venus_"):
                                    data.loc[key[0]+col, '适用车型'] = "[\"Venus\"]"
                                else:
                                    data.loc[key[0]+col, '适用车型'] = "[\"Mars One\"]"
                                
            else:
                # 插入用例名称，步骤和预期结果
                data.loc[key[0]+col, '用例名称'] = case_name
                data.loc[key[0]+col, '步骤描述'] = step_name
                data.loc[key[0]+col, '预期结果'] = f"【1】同帧非相关信号跳变无异常{interface_name}事件"
                # 其他列填入默认值
                data.loc[key[0]+col, '所属模块'] = "/SOA服务接口/通用场景/信号跳变校验event事件场景/AllService"
                data.loc[key[0]+col, '整车版本'] = "[\"V2.0.0 Beta2\"]"
                data.loc[key[0]+col, '编辑模式'] = "STEP"
                data.loc[key[0]+col, '标签'] = "two_domain"
                data.loc[key[0]+col, '用例类型'] = "接口用例"
                data.loc[key[0]+col, '责任人'] = "陶磊 Lei(SOA)"
                data.loc[key[0]+col, '用例等级'] = "P3"
                data.loc[key[0]+col, '适用范围'] = "Full"
                data.loc[key[0]+col, '业务级别'] = "[PlatformTest-SOA]"
                data.loc[key[0]+col, 'Automation'] = "Automated"
                data.loc[key[0]+col, '测试平台'] = "Bench"
                data.loc[key[0]+col, "函数名称"] = interface_name
                if sdbsignals_sheet.startswith("Venus_"):
                    data.loc[key[0]+col, '适用车型'] = "[\"Venus\"]"
                else:
                    data.loc[key[0]+col, '适用车型'] = "[\"Mars One\"]"

        data.to_excel(ms_excel_path, index=False)


def write_template(data, id=None, service_name=None, interface_name=None, network_name=None, message_ID=None, signal_list=None, case_name=None, car_type=None):
    file = os.path.join(current_path, "test_signal_jump_check_event.py")

    if not os.path.exists(file):
        with open(file, "w") as f:
            pass

    with open(file, "a") as fd:
        for i in data:
            if '{ID}' in i:
                fd.write(i.replace('{ID}', id))
            elif '{service}' in i:
                if car_type == "[\"Venus\"]":
                    fd.write(i.replace('{service}', service_name + 'Venus'))
                else:
                    fd.write(i.replace('{service}', service_name))
            elif '{service_name}' in i:
                fd.write(i.replace('{service_name}', '"' + service_name + '"'))
            elif 'udp_mcu_ip' in i:
                fd.write(i)
                if car_type == "[\"Venus\"]":
                    fd.write('        self.sd_tester.write_single_ccp(950, 2)' + '\n')
            elif '{interface_name}' and '{eventorreq}' in i:
                if "ACCService" in service_name:
                    fd.write(i.replace('{eventorreq}', 'req').replace('{interface_name}', '"' + interface_name + '"'))
                else:
                    fd.write(i.replace('{eventorreq}', 'event').replace('{interface_name}', '"' + interface_name + '"'))
            elif '{network_name}' and '{message_ID}' and '{signal_list}' in i:
                if '0x' in message_ID:
                    msg_id = int(message_ID, 16)
                else:
                    msg_id = int(message_ID.split('-')[0])*256*256 + int(message_ID.split('-')[1])*256 + int(message_ID.split('-')[-1])
                # 有UB位控制的信号会干扰event发送，添加判断逻辑
                ub_flag=0
                for i in signal_list:
                    if "_UB" in i:
                        ub_flag=1
                        break
                if ub_flag==0:
                    fd.write(i.replace('{message_ID}', str(msg_id)).replace('{network_name}', '"' + network_name + '"').replace('{signal_list}', str(signal_list))).replace('{UB_Flag}', '')
                else:
                    fd.write(i.replace('{message_ID}', str(msg_id)).replace('{network_name}', '"' + network_name + '"').replace('{signal_list}', str(signal_list))).replace('{UB_Flag}', ', UB_Flag=False')
            elif '{case_name}' in i:
                fd.write(i.replace('{case_name}', '"' + case_name + '"'))
            else:
                fd.write(i)


def put_case_into_python(ms_file):
    """
    把ms用例输出成python脚本到test_signal_jump_check_event.py
    需要下载ms上带caseid的excel表格
    """
    # 生成所有适用mars1车型的脚本
    case_info = pd.read_excel(os.path.join(current_path.split("sat")[0], ms_file), sheet_name='模版', keep_default_na=True)
    case_data_dict = case_info.to_dict()
    service_name_begin = ''

    with open(os.path.join(current_path, "update_template.txt"), "r") as fd:
        lines = fd.readlines()
        template_from = lines[:14]
        template_class = lines[14:35]
        template_method = lines[35:] + ['\n']
        write_template(template_from)
        for i in range(len(case_data_dict['ID'])):
            id = str(case_data_dict['ID'][i])
            service_name = case_data_dict['用例名称'][i].split('::')[0]
            interface_name = case_data_dict['用例名称'][i].split('::')[1].split('_')[0]
            network = case_data_dict['用例名称'][i].split('::')[1].split('_')
            if len(network) == 2:
                network_name = network[1].lower()
            elif len(network) == 3:
                network_name = network[1].lower() + '_' + network[-1].lower()
            message_ID = case_data_dict['用例名称'][i].split('::')[2].split('同帧')[0]
            signal_list = case_data_dict['步骤描述'][i].split('中')[-1].split('以外的')[0].split('，')
            car_type = case_data_dict['适用车型'][i]
            if service_name != service_name_begin:
                write_template(template_class, id, service_name, interface_name, network_name, message_ID, signal_list, case_data_dict['用例名称'][i], car_type)
            service_name_begin = service_name
            write_template(template_method, id, service_name, interface_name, network_name, message_ID, signal_list, case_data_dict['用例名称'][i], car_type)

    # 生成只适用Venus车型的脚本
    # case_info_venus = pd.read_excel(os.path.join(current_path.split("sat")[0], ms_file), sheet_name='Venus', keep_default_na=True)
    # case_data_dict_venus = case_info_venus.to_dict()

    # with open(os.path.join(current_path, "update_template.txt"), "r") as fd:
    #     lines = fd.readlines()
    #     template_from = lines[:14]
    #     template_class = lines[14:34]
    #     template_method = lines[34:] + ['\n']
    #     for i in range(len(case_data_dict_venus['ID'])):
    #         id = str(case_data_dict_venus['ID'][i])
    #         service_name = case_data_dict_venus['用例名称'][i].split('::')[0]
    #         interface_name = case_data_dict_venus['用例名称'][i].split('::')[1].split('_')[0]
    #         network_name = case_data_dict_venus['用例名称'][i].split('::')[1].split('_')[1].lower()
    #         message_ID = case_data_dict_venus['用例名称'][i].split('::')[2].split('同帧')[0]
    #         signal_list = case_data_dict_venus['步骤描述'][i].split('中')[-1].split('以外的')[0].split('，')
    #         car_type = case_data_dict_venus['适用车型'][i]
    #         if service_name != service_name_begin:
    #             write_template(template_class, id, service_name, interface_name, network_name, message_ID, signal_list, case_data_dict_venus['用例名称'][i], car_type)
    #         service_name_begin = service_name
    #         write_template(template_method, id, service_name, interface_name, network_name, message_ID, signal_list, case_data_dict_venus['用例名称'][i], car_type)


def change_intercasename_reqid(from_file, ms_file):
    # 获取需要修改REQ和接口名称的case信息
    case_info = pd.read_excel(os.path.join(current_path.split("sat")[0], ms_file), keep_default_na=True)
    case_data_dict = case_info.to_dict()
    # 获取所有ms上接口中的case信息
    from_dict = {}
    for file in os.listdir(os.path.join(current_path.split("sat")[0], from_file)):
        print(f"开始更新文件{file}")
        from_case_info = pd.read_excel(os.path.join(current_path.split("sat")[0], from_file, file), keep_default_na=True)
        from_dict.update(from_case_info.to_dict()) 
    # 开始修改REQ和接口名称的case信息
    for i in range(len(case_data_dict['用例名称'])):
        for j in range(len(from_dict['ID'])):
            
            if case_data_dict['函数名称'][i] == from_dict['函数名称'][j]:
                case_info.loc[i, 'JamaReqID'] = from_dict['JamaReqID'][j]
                case_info.loc[i, '接口名称'] = from_dict['接口名称'][j]
            elif case_data_dict['函数名称'][i] == from_dict['接口名称'][j]:
                case_info.loc[i, 'JamaReqID'] = from_dict['JamaReqID'][j]
                case_info.loc[i, '接口名称'] = from_dict['函数名称'][j]
            elif case_data_dict['函数名称'][i] in str(from_dict['函数名称'][j]).replace('，', ',').split(','):
                print(from_dict['函数名称'][j].replace('，', ',').split(','))
                case_info.loc[i, 'JamaReqID'] = from_dict['JamaReqID'][j]
                case_info.loc[i, '接口名称'] = from_dict['接口名称'][j]
            elif case_data_dict['函数名称'][i] in str(from_dict['接口名称'][j]).replace('，', ',').split(','):
                case_info.loc[i, 'JamaReqID'] = from_dict['JamaReqID'][j]
                case_info.loc[i, '接口名称'] = from_dict['接口名称'][j]

    case_info.to_excel(os.path.join(current_path.split("sat")[0], ms_file), index=False)



if __name__ == "__main__":
    # put_caseinfo_into_ms("JIDL24R1_Baseline20240607_SDBSignals_20240520_Release.xlsx","Metersphere_case_SOA平台测试.xlsx")
    # 需要分开跑，下面需上传ms后获取caseid再下载下来生成脚本
    put_case_into_python("caseid_Metersphere_case_SOA平台测试.xlsx")
