#!/usr/bin/env python 3.8
# _*_ coding: utf-8 _*_

"""
@Time: 2023/10/23 13:23
@Author: lei.tao
@File: update_ms_case_info.py
@Software: PyCharm
@Description: 更新ms用例信息
@Example:
"""
import os
import sys
import pandas as pd
import re
import openpyxl
project_root = os.path.join(os.getcwd().split("sat")[0], 'sat')
sys.path.append(project_root)


def has_chinese(string):
    pattern = re.compile(r'[\u4e00-\u9fa5]')
    match = pattern.search(string)
    return match


def is_english(string):
    pattern = "^[A-Za-z0-9,]+$"
    if re.match(pattern, string):
        return True
    else:
        return False


def to_new_excel(old_file, new_file):
    """
    修复ms用例的信息，生成新的excel表格
    """
    wb = openpyxl.load_workbook(old_file)
    sheet = wb.active
    header = []
    if not header:
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(1, c).value
            if val:
                header.append(val)

    for row in sheet.iter_rows(min_row=2, min_col=1, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            if cell.value:
                cell.value = str(cell.value).replace("null", ",").replace("，", ",").replace('"', '').replace('“', '').replace('”', '').replace('[]', '')
    wb.save(filename=new_file)
    # 添加一列：修改意见
    if "修改意见" not in header:
        sheet.insert_cols(1, 1)
        sheet.cell(1, 1, value="修改意见")

    wb.save(filename=new_file)
    # 修改ms信息
    data = pd.read_excel(new_file, keep_default_na=False)
    # data.fillna(method='ffill', inplace=True)
    print(data.to_dict())
    for i in range(len(data.to_dict()['ID'])):
        if data.to_dict()['ID'][i] != '':
            # 校验函数名称和接口名称
            if not is_english(data.to_dict()['函数名称'][i]):
                if is_english(data.to_dict()['接口名称'][i]):
                    cell_value_func = data.to_dict()['函数名称'][i]
                    cell_value_interface = data.to_dict()['接口名称'][i]
                    data.loc[i, '函数名称'] = cell_value_interface
                    data.loc[i, '接口名称'] = cell_value_func
                else:
                    data.loc[i, '修改意见'] = "函数名称需要修改"
            else:
                if is_english(data.to_dict()['接口名称'][i]):
                    data.loc[i, '修改意见'] = "接口名称需要修改"
            # 把需求ID信息copy到JamaReqID
            if data.to_dict()['需求ID'][i] == '':
                data.loc[i, '修改意见'] = "需求ID不能为空，请填写需求ID和JamaReqID"
            else:
                data.loc[i, 'JamaReqID'] = data.to_dict()['需求ID'][i]
            # 对于空的地段填写默认值
            if data.to_dict()['用例类型'][i] == '':
                data.loc[i, '用例类型'] = "接口用例"
            if data.to_dict()['适用车型'][i] == '[]' or data.to_dict()['适用车型'][i] == '':
                data.loc[i, '适用车型'] = "[Mars One]"
            if data.to_dict()['用例等级'][i] == '':
                data.loc[i, '用例等级'] = "P2"
            if data.to_dict()['适用范围'][i] == '':
                data.loc[i, '适用范围'] = "Full"
            if data.to_dict()['业务级别'][i] == '':
                data.loc[i, '业务级别'] = "[PlatformTest-SOA]"
            if data.to_dict()['Automation'][i] != 'Automated':
                data.loc[i, 'Automation'] == "Automated"
            if data.to_dict()['整车版本'][i] == '[]' or data.to_dict()['整车版本'][i] == '':
                data.loc[i, '整车版本'] = "[1.1]"
            if data.to_dict()['测试平台'][i] == '':
                data.loc[i, '测试平台'] = "Bench"
        else:
            continue

    data.to_excel(new_file, index=False)


def _merge_cell(ws, target_list, start_row, col):
    """
    合并重复的单元格
    """
    start = 0
    end = 0
    ref = target_list[0]
    lists = []
    for i in range(len(target_list)):
        if target_list[i] != ref:
            ref = target_list[i]
            end = i - 1
            ws.merge_cells(col + str(start + start_row) + ":" + col + str(end + start_row))
            start = end + 1
        if i == len(target_list) - 1:
            end = i
            ws.merge_cells(col + str(start + start_row) + ":" + col + str(end + start_row))


def get_cell_loc(ws):
    """
    获取需要合并的空单元格位置
    """
    start = 0
    end = 0
    target_list = []
    for row in range(2, ws.max_row + 1):
        id = ws['B' + str(row)].value
        target_list.append(id)
    lists = []
    for i in range(len(target_list) - 1):
        if target_list[i] != None:
            if target_list[i + 1] == None:
                lists.append({start: end})
                end = 0
                start = i + 2
                end += 1
        else:
            end += 1

    if target_list[-1] == None:
        end += 1
    lists.append({start: end})

    return lists[1:]


def merge_cell(new_file):
    """
    合并单元格
    """
    wb = openpyxl.load_workbook(new_file)
    ws = wb['Sheet1']
    data = get_cell_loc(ws)
    title = []
    for letter in range(65, 89):
        title.append(chr(letter).upper())
    del title[6]
    del title[6]
    print(title)
    for a in range(len(title)):
        list_id = []
        for row in range(2, ws.max_row + 1):
            id = ws[title[a] + str(row)].value
            list_id.append(id)
        print(list_id)
        for i in data:
            ws.merge_cells(
                title[a] + str(list(i.keys())[0]) + ":" + title[a] + str(list(i.keys())[0] + list(i.values())[0] - 1))

    wb.save(new_file)
    wb.close()


if __name__ == '__main__':
    path = "C:\\Users\\lei.tao\\Downloads\\Metersphere_case_SOA_3"
    file_list = os.listdir(path)
    print(file_list)
    for file in file_list:
        file_path = os.path.join(path, file)
        new_file_path = os.path.join(path, f"new_{file}")
        to_new_excel(file_path, new_file_path)
        merge_cell(new_file_path)
