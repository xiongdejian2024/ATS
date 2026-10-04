# -*- coding: utf-8 -*-
"""
@File        : test_all_in1.py
@Author      : jinhui.zhao@jiduauto.com
@Time        : 2022/7/17 21:25
@Description : 
@Examples    :
"""

import openpyxl
from openpyxl.styles import Alignment
from xat_ecu.legacy.common.logger import logger

sheet_list = ['ChassisService', 'TailGateService', 'CentralLockService', 'BonnetService', 'ChargeLidService', 'ClimateControl', 'DoorService', 'VehicleModeService', 'TyreService', 'DrivingAssistService', 'GloveBoxService', 'HornService', 'InnerRearViewService', 'KeyService', 'LightService', 'OuterRearViewService', 'PassiveSafetyService', 'LowVoltageService', 'PedalService', 'SeatService', 'ShieldWindowService', 'SteerWheelService', 'SunroofService']





if __name__ == "__main__":
    case_file = "S2S测试用例_v055.xlsx"
    wb = openpyxl.load_workbook(case_file)
    ws_n = wb.create_sheet("111")
    n_row = 1
    header = []
    for sn in wb.sheetnames:
        if sn in sheet_list:
            sheet = wb.get_sheet_by_name(sn)
            logger.info(f'{sn} in list')
            if header == []:
                logger.info(f'{sn} for header')
                for c in range(1, sheet.max_column + 1):  # header
                    val = sheet.cell(1, c).value
                    if val:
                        header.append(val)
                        ws_n.cell(row=n_row, column=c, value=val)
            # case
            for r in range(2, sheet.max_row+1):  # 从第二行开始
                # 去除首个单元格为空的行
                if sheet.cell(r, 1).value == None:
                    # logger.info(f'{sn} row: {r}, colum: {1} skipped')
                    continue
                if isinstance(sheet.cell(r, 1).value, str):
                    if sheet.cell(r, 1).value.strip() == "":
                        # logger.info(f'{sn} row: {r}, colum: {1} skipped')
                        continue
                n_row += 1
                for c in range(1, sheet.max_column+1):
                    val = sheet.cell(r, c).value
                    ws_n.cell(row=n_row, column=c, value=val).alignment = Alignment(wrap_text=True)
    wb.save(filename="ffff.xlsx")




