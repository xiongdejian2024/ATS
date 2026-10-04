# -*- coding: utf-8 -*-
"""
@File        : statistics_tool.py
@Author      : quan.sun@jiduauto.com
@Time        : 2023/01/03 12:18
@Description :
@Examples    :
"""

import os
import sys
current_path = os.path.dirname(os.path.realpath(__file__))
sys.path.append(current_path)
sys.path.append(os.path.join(current_path, ".."))
sys.path.append(os.path.join(current_path, "../.."))
sys.path.append(os.path.join(current_path, "../../.."))
import openpyxl
from openpyxl import *
from xat_ecu.legacy.common.logger import *


class StatisticsTool:
    def __init__(self, case_file_path, statistics_file_path):
        self.case_file_path = case_file_path
        self.statistics_file_path = statistics_file_path

    def get_cases_data(self):
        wb = load_workbook(self.case_file_path)
        ws = wb.active

        for column in ws.iter_cols():
            if column[0].value == "所属模块":
                module = column[1:]
            elif column[0].value == "Automation":
                Automatable = column[1:]

        datas = []
        length = len(module)
        for i in range(0, length):
            data = []
            module_name = module[i].value
            module_name_list = module_name.split("/")
            module_name1 = module_name_list[1]
            data.append(module_name1)
            module_name2 = module_name_list[2]
            if module_name2 == "网络架构":
                module_name2 = module_name2 + "_" + module_name_list[3]
            data.append(module_name2)

            if Automatable[i].value == "Ready":
                Automated_value = "Y"
                Automatable_value = "Y"
            elif Automatable[i].value == "To Do":
                Automated_value = "N"
                Automatable_value = "Y"
            else:
                Automated_value = "N"
                Automatable_value = "N"

            data.append(Automatable_value)
            data.append(Automated_value)

            datas.append(data)
        wb.close()
        return datas

    def write_data_into_excel(self):
        wb = Workbook()
        ws = wb.active

        ws['A1'] = "序号"
        ws['B1'] = "一级模块"
        ws['C1'] = "二级模块"
        ws['D1'] = "用例数"
        ws['E1'] = "Automatable"
        ws['F1'] = "Automated"
        ws['G1'] = "Automatable Rate"
        ws['H1'] = "Automated Rate"
        ws['I1'] = "计划新增自动化case数"
        ws['J1'] = "实际完成自动化case数"
        ws['K1'] = "备注"

        datas = self.get_cases_data()
        rows = {}
        for data in datas:
            module_name1 = data[0]
            module_name2 = data[1]
            if  rows.get(module_name1) is None:
                rows[module_name1] = {}
            if rows[module_name1].get(module_name2) is None:
                rows[module_name1][module_name2] = [0, 0, 0]     # "用例数","Automatable","Automated"
            row = rows[module_name1][module_name2]
            row[0] += 1
            if data[2] == "Y":
                row[1] += 1
            if data[3] == "Y":
                row[2] += 1

        i = 1
        for name1 in rows:
            for name2 in rows[name1]:
                row_list = rows[name1][name2]
                case = row_list[0]
                Automatable = row_list[1]
                Automated = row_list[2]
                Automatable_Rate = Automatable/case
                if Automatable == 0:
                    Automated_Rate = 0
                else:
                    Automated_Rate = Automated/Automatable

                i += 1
                ws.cell(i, 1).value = i - 1
                ws.cell(i, 2).value = name1
                ws.cell(i, 3).value = name2
                ws.cell(i, 4).value = case
                ws.cell(i, 5).value = Automatable
                ws.cell(i, 6).value = Automated
                ws.cell(i, 7).value = Automatable_Rate
                ws.cell(i, 7).number_format = '0.00%'
                ws.cell(i, 8).value = Automated_Rate
                ws.cell(i, 8).number_format = '0.00%'
                logger.info(name2)

        wb.save(self.statistics_file_path)
        wb.close()



if __name__ == "__main__":
    # Work Path: sat/
    # Command: python3 tools/case_tools/statistics_tool.py
    logger = Logger().get_logger("test")

    xmindtool = StatisticsTool("/root/quansun/sat/tools/case_tools/Metersphere_case_BGM_ComponentTest_MarsOne.xlsx", "/root/quansun/sat/tools/case_tools/bgm_out.xlsx")
    xmindtool.write_data_into_excel()

    xmindtool = StatisticsTool("/root/quansun/sat/tools/case_tools/Metersphere_case_TCAM_ComponentTest_MarsOne.xlsx", "/root/quansun/sat/tools/case_tools/tcam_out.xlsx")
    xmindtool.write_data_into_excel()