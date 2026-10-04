import os
import sys
import openpyxl
import pandas as pd

project_root = __import__("xat_ecu.resources", fromlist=["LEGACY_ROOT"]).LEGACY_ROOT.as_posix()
from xat_ecu.legacy.common.logger import logger


def get_caseinfo():
    base = os.path.join(project_root, "test_case/soa/test_function")
    case_info = {k.split("_")[1].replace(".py", "").lower(): [] for k in os.listdir(base)}
    class_mod = ""
    for file in os.listdir(base):
        k = file.split("_")[1].replace(".py", "").lower()
        if file.startswith("test_") and file.endswith(".py"):
            logger.info(f"开始处理文件{file}")
            with open(os.path.join(base, file), "r", encoding="utf8") as fd:
                lines = fd.read().split("\n")
                class_all = []
                for j in range(len(lines)):
                    if "class " in lines[j] and "(TestBase" in lines[j]:
                        class_all.append((j, lines[j].split("class ")[-1].split("(TestBase")[0]))

                if len(class_all) == 0:
                    assert False, f"{file}文件用例类查询失败，请先验证"

                for i in range(len(lines)):
                    if "caseid" in lines[i] and "(self)" in lines[i]:
                        id = lines[i].split("_caseid_")[-1].split("(self)")[0]
                        module = os.path.join("test_case/soa/test_function", file).replace("\\", "/")
                        func = lines[i].split("def ")[-1].split("(self)")[0].strip()

                        if len(class_all) == 1:
                            class_mod = class_all[0][-1]

                        else:
                            for a in range(0, (len(class_all) - 1)):
                                if class_all[a][0] < i < class_all[a + 1][0]:
                                    class_mod = class_all[a][-1]
                                else:
                                    class_mod = class_all[a + 1][-1]
                        if "_" in id:
                            for j in id.split("_"):
                                case_info[k].append((j, [{"num": 1, "module": module, "class": class_mod, "func": func}]))
                        else:
                            case_info[k].append((id, [{"num": 1, "module": module, "class": class_mod, "func": func}]))
        else:
            continue

    return case_info


def to_new_excel(old_file, new_file):
    """
    修复ms用例的信息，生成新的excel表格
    """
    wb = openpyxl.load_workbook(old_file)
    sheet = wb.active
    header = []
    if not header:
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(1, c).value
            if val:
                header.append(val)

    wb.save(filename=new_file)
    # 修改ms信息
    data = pd.read_excel(new_file, keep_default_na=False)
    caseinfo = get_caseinfo()
    logger.info(caseinfo)
    for i in range(len(data.to_dict()['ID'])):
        if data.to_dict()['ID'][i] != '':
            data.loc[i, '自动化用例代码仓库'] = "https://jidudev.com/soa/soa_test/sat.git"
            k = data.to_dict()['所属模块'][i].split("/")[3].lower()
            if k in caseinfo.keys():
                for j in range(len(caseinfo[k])):
                    if str(data.to_dict()['ID'][i]) == caseinfo[k][j][0]:
                        info = caseinfo[k][j][1]
                        logger.info(info)
                        data.loc[i, '自动化用例'] = info
                        break
                else:
                    data.loc[i, '自动化用例'] = '[{"num": 1, "module": "", "class": "", "func": ""}]'
            else:
                pass

        data.to_excel(new_file, index=False)


def get_cell_loc(ws):
    """
    获取需要合并的空单元格位置
    """
    start = 0
    end = 0
    target_list = []
    for row in range(2, ws.max_row + 1):
        id = ws['A' + str(row)].value
        target_list.append(id)
    lists = []
    for i in range(len(target_list) - 1):
        if target_list[i] is not None:
            if target_list[i + 1] is None:
                lists.append({start: end})
                end = 0
                start = i + 2
                end += 1
        else:
            end += 1

    if target_list[-1] is None:
        end += 1
    lists.append({start: end})

    return lists[1:]


def merge_cell(new_file):
    """
    合并单元格
    """
    wb = openpyxl.load_workbook(new_file)
    ws = wb['Sheet1']
    data = get_cell_loc(ws)
    title = []
    for letter in range(65, 89):
        title.append(chr(letter).upper())
    # F,G两列不合并
    del title[5]
    del title[5]
    for a in range(len(title)):
        list_id = []
        for row in range(2, ws.max_row + 1):
            id = ws[title[a] + str(row)].value
            list_id.append(id)
        for i in data:
            ws.merge_cells(
                title[a] + str(list(i.keys())[0]) + ":" + title[a] + str(list(i.keys())[0] + list(i.values())[0] - 1))

    wb.save(new_file)
    wb.close()


if __name__ == '__main__':
    path = "C:\\Users\\Downloads\\111111111111111111111"
    file_list = os.listdir(path)
    logger.info(file_list)
    for file in file_list:
        file_path = os.path.join(path, file)
        new_file_path = os.path.join(path, f"new_{file}")
        to_new_excel(file_path, new_file_path)
        merge_cell(new_file_path)
