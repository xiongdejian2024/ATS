# -*- coding: utf-8 -*-
"""
@File        : statistics_data.py
@Author      : quan.sun@jiduauto.com
@Time        : 2023/05/30 12:18
@Description :
@Examples    :
"""

import os
import sys

current_path = os.path.dirname(os.path.realpath(__file__))
import openpyxl
from openpyxl import *
from xat_ecu.legacy.common.logger import *
from xat_ecu.legacy.interface.ms.ms_lib import meterSphere_client


'''  ======= test plan case info  ========
{'id': 'da4863cc-946f-4e8e-8c7e-05f7ab40e52e', 
'nodeId': 'd56185f7-d733-48b1-bc8e-897af92d03cd',
'testId': '[]', 
'nodePath': '/架构基础/数字安全/SOA服务鉴权',
'projectId': 'f1d8bad3-4417-4e02-8b8d-28e583cd2e80', 
'name': 'SOA服务鉴权_服务启动service_monitor', 
'type': 'functional', 
'maintainer': 'keke.wu', 
'priority': 'P0', 
'method': '', 
'createTime': 1680694602353,
'updateTime': 1685503835521, 
'sort': None, 
'num': 1567664, 
'otherTestName': None, 
'reviewStatus': None, 
'tags': '["smoke"]', 
'demandId': None, 
'demandName': None, 
'status': 'Pass',
'stepModel': None,
'customNum': '1567664', 
'createUser': None,
'originalStatus': None, 
'deleteTime': None, 
'deleteUserId': None, 
'order': None, 
'casePublic': None,
'versionId': 'b7290d69-db5e-4115-abb7-ee9a67ea26c4',
'refId': None, 
'latest': None, 
'prerequisite': None,
'remark': None, 
'steps': None, 
'stepDescription': None,
'expectedResult': None, 
'customFields': '[{"id":"1c93d5bd-ee73-4200-b464-207d402279e6",
                    "name":"JamaProjectId","value":"MSO-TC-34820"},
                {"id":"5e0e5b6c-bc72-4c74-9b97-c4ccd76c4e9c",
                    "name":"JamaGlobalId",
                    "value":"GID-528268"},
                {"id":"e728890c-410b-413f-972e-fdb4c2a0e7ab",
                    "name":"TestExecutionMode","value":"Automation"},
                {"id":"bcbdeecf-4978-4463-879c-00252642715e",
                    "name":"Automation",
                    "value":"Ready"},
                {"id":"150a702a-a5a1-493e-bca1-c567665328ac",
                    "name":"AffectsVersion",
                    "value":[]},
                {"id":"98df5b85-e534-4489-b7f9-d41215eeaaf1",
                    "name":"责任人",
                    "value":"keke.wu"},
                {"id":"a71341de-de5b-4775-8266-cd1d13753471",
                    "name":"用例等级",
                    "value":"P0"},
                {"id":"acfae654-8107-4b1c-b29a-53bb290be66b",
                    "name":"用例状态",
                    "value":"Prepare"},
                {"id":"04202723-5e33-42a6-b250-aa4ac14c753a",
                    "name":"Test_Level",
                    "value":"a5ed29d2"}]',
'executor': 'ldap_soa_jama', 
'executorName': 'ldap_soa_jama', 
'results': None, 
'planId': '1bd1ba6f-f9fd-42ea-a45d-93ab457e4d0b', 
'planName': None,
'caseId': '0ac77181-9a0c-f032-4446-3702e1e17b2e', 
'issues': None, 
'reportId': None, 
'model': None, 
'projectName': 'TCAM_ComponentTest_MarsOne', 
'actualResult': '', 
'maintainerName': '吴柯柯 Keke', 
'isCustomNum': None,
'issuesCount': 0, 
'versionName': None, 
'outerLink': None, 
'list': None, 
'issueList': None}
'''


class StatisticsData:
    def __init__(self, case_file_path, statistics_file_path, testplan):
        self.case_file_path = case_file_path
        self.statistics_file_path = statistics_file_path
        self.testplan = testplan

    def get_cases_datas(self):
        # 获取并转换计算数据
        wb = load_workbook(self.case_file_path)
        ws = wb.active

        client = meterSphere_client()
        report_info = client.get_report_info_from_planid(self.testplan)
        # logger.info(report_info)
        for sectitle in report_info.keys():
            case_info = report_info[sectitle]
            case_info["用例数"] = len(case_info["一级模块"])
            case_info["一级模块"] = case_info["一级模块"][0]

            # "96ae3bb3" ---- Ready
            # "06ca513d" ---- To Do
            case_info["Automatable"] = case_info["Automation"].count(
                "96ae3bb3"
            ) + case_info["Automation"].count("06ca513d")
            if case_info["Automatable"] == 0:
                case_info["Automatable"] = case_info["Automation"].count(
                    "Ready"
                ) + case_info["Automation"].count("To Do")
            case_info["Pass"] = case_info["status"].count("Pass")
            case_info["Failure"] = case_info["status"].count("Failure")
            case_info["Error"] = case_info["status"].count("Error")
            case_info["Automated"] = (
                case_info["Pass"] + case_info["Failure"] + case_info["Error"]
            )
            if case_info["用例数"]:
                case_info["Automatable Rate"] = (
                    case_info["Automatable"] / case_info["用例数"]
                )
            if case_info["Automatable"]:
                case_info["Automated Rate"] = (
                    case_info["Automated"] / case_info["Automatable"]
                )
            else:
                case_info["Automated Rate"] = 0

        for column in ws.iter_cols():
            # logger.info(column[0].value)
            if column[0].value == "自动化Owner":
                a_owners = column[1:]
            elif column[0].value == "功能Owner":
                f_owners = column[1:]
            elif column[0].value == "一级模块":
                f_modules = column[1:]
            elif column[0].value == "二级标题":
                sec_titles = column[1:]
            elif column[0].value == "备注":
                notes = column[1:]
            elif column[0].value == "是否ok":
                is_oks = column[1:]
            elif column[0].value == "Block":
                blocks = column[1:]
            elif column[0].value == "风险":
                risks = column[1:]

        sec_titles_not_match_list = []
        soa_sec_title = ""
        for i in range(len(column[1:])):
            if f_modules[i].value and sec_titles[i].value:
                soa_sec_title = f_modules[i].value + "_" + sec_titles[i].value
            else:
                soa_sec_title = ""
            if sec_titles[i].value in report_info:
                report_info[sec_titles[i].value]["自动化Owner"] = a_owners[i].value
                report_info[sec_titles[i].value]["功能Owner"] = f_owners[i].value
                report_info[sec_titles[i].value]["备注"] = notes[i].value
                report_info[sec_titles[i].value]["是否ok"] = is_oks[i].value
                report_info[sec_titles[i].value]["Block"] = blocks[i].value
                report_info[sec_titles[i].value]["风险"] = risks[i].value
            elif soa_sec_title in report_info:  # 说明是SOA
                report_info[soa_sec_title]["自动化Owner"] = a_owners[i].value
                report_info[soa_sec_title]["功能Owner"] = f_owners[i].value
                report_info[soa_sec_title]["备注"] = notes[i].value
                report_info[soa_sec_title]["是否ok"] = is_oks[i].value
                report_info[soa_sec_title]["Block"] = blocks[i].value
                report_info[soa_sec_title]["风险"] = risks[i].value
            elif sec_titles[i].value == "Total":
                pass
            else:
                sec_titles_not_match_list.append(sec_titles[i].value)
        logger.info("sec_titles not match list : {}".format(sec_titles_not_match_list))

        wb.close()
        return report_info

    def write_data_into_excel(self):
        wb = Workbook()
        ws = wb.active

        ws['A1'] = "序号"
        ws['B1'] = "一级模块"
        ws['C1'] = "二级标题"
        ws['D1'] = "用例数"
        ws['E1'] = "Automatable"
        ws['F1'] = "Automated"
        ws['G1'] = "Automatable Rate"
        ws['H1'] = "Automated Rate"
        ws['I1'] = "自动化Owner"
        ws['J1'] = "功能Owner"
        ws['K1'] = "Pass"
        ws['L1'] = "Failure"
        ws['M1'] = "Error"  #  MS 里的 Blocking

        # 附加
        ws['N1'] = "备注"
        ws['O1'] = "是否ok"
        ws['P1'] = "Block"
        ws['Q1'] = "风险"

        report_info = self.get_cases_datas()
        logger.info("========== 开始写入 Case =========")
        i = 1
        case_total = 0
        Automatable_total = 0
        Automated_total = 0
        Pass_total = 0
        Failure_total = 0
        Error_total = 0
        block_total = 0
        risk_total = 0
        for sec_title, case_info in report_info.items():
            # logger.info(sec_title)
            # logger.info(case_info)
            i += 1
            ws.cell(i, 1).value = i - 1

            if case_info.get("一级模块") == "SOA服务接口":
                ws.cell(i, 2).value = sec_title.split("_")[0]
                ws.cell(i, 3).value = sec_title.split("_")[-1]
            else:
                ws.cell(i, 2).value = case_info.get("一级模块")
                ws.cell(i, 3).value = sec_title

            ws.cell(i, 4).value = case_info.get("用例数")
            case_total += case_info.get("用例数")
            ws.cell(i, 5).value = case_info.get("Automatable")
            Automatable_total += case_info.get("Automatable")
            ws.cell(i, 6).value = case_info.get("Automated")
            Automated_total += case_info.get("Automated")
            ws.cell(i, 7).value = case_info.get("Automatable Rate")
            ws.cell(i, 7).number_format = '0.00%'
            ws.cell(i, 8).value = case_info.get("Automated Rate")
            ws.cell(i, 8).number_format = '0.00%'

            ws.cell(i, 9).value = case_info.get("自动化Owner")
            ws.cell(i, 10).value = case_info.get("功能Owner")
            ws.cell(i, 11).value = case_info.get("Pass")
            Pass_total += case_info.get("Pass")
            ws.cell(i, 12).value = case_info.get("Failure")
            Failure_total += case_info.get("Failure")
            ws.cell(i, 13).value = case_info.get("Error")
            Error_total += case_info.get("Error")

            ws.cell(i, 14).value = case_info.get("备注")
            ws.cell(i, 15).value = case_info.get("是否ok")
            ws.cell(i, 16).value = case_info.get("Block")
            if case_info.get("Block"):
                block_total += case_info.get("Block")
            ws.cell(i, 17).value = case_info.get("风险")
            if case_info.get("风险"):
                risk_total += case_info.get("风险")

            logger.info(sec_title)

        ws.cell(i + 1, 3).value = "Total"
        ws.cell(i + 1, 4).value = case_total
        ws.cell(i + 1, 5).value = Automatable_total
        ws.cell(i + 1, 6).value = Automated_total
        if case_total:
            ws.cell(i + 1, 7).value = Automatable_total / case_total
        ws.cell(i + 1, 7).number_format = '0.00%'
        if Automatable_total:
            ws.cell(i + 1, 8).value = Automated_total / Automatable_total
        ws.cell(i + 1, 8).number_format = '0.00%'

        ws.cell(i + 1, 11).value = Pass_total
        ws.cell(i + 1, 12).value = Failure_total
        ws.cell(i + 1, 13).value = Error_total

        ws.cell(i + 1, 16).value = block_total
        ws.cell(i + 1, 17).value = risk_total

        wb.save(self.statistics_file_path)
        wb.close()
        logger.info("========== 写入 Case 结束 =========")


if __name__ == "__main__":
    # Work Path: sat/
    # Command: python3  ecu_simulator/interface/ms/statistics_data.py --bgm_id="ccbfb89c-68bc-4da2-bc76-548794d5fce3"

    import argparse

    argparser = argparse.ArgumentParser()
    argparser.add_argument(
        '--bgm_id',
        help='test plan id of bgm',
        default="ccbfb89c-68bc-4da2-bc76-548794d5fce3",
    )
    argparser.add_argument(
        '--tcam_id',
        help='test plan id of tcam',
        default="1bd1ba6f-f9fd-42ea-a45d-93ab457e4d0b",
    )
    argparser.add_argument(
        '--soa_id',
        help='test plan id of soa',
        default="bc2c30af-f731-4e16-8a2d-0d191e9eb01b",
    )
    # Default is All
    args = argparser.parse_args()

    logger = Logger().get_logger("test")

    bgm = StatisticsData(
        "/root/quansun/bgm_2023_06_12.xlsx",
        "/root/quansun/bgm_out.xlsx",
        args.bgm_id,
    )
    bgm.write_data_into_excel()

    tcam = StatisticsData(
        "/root/quansun/tcam_2023_06_12.xlsx",
        "/root/quansun/tcam_out.xlsx",
        args.tcam_id,
    )
    tcam.write_data_into_excel()

    soa = StatisticsData(
        "/root/quansun/soa_2023_06_12.xlsx",
        "/root/quansun/soa_out.xlsx",
        args.soa_id,
    )
    soa.write_data_into_excel()
